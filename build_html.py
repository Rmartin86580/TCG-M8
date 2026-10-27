import json, os
from openpyxl import load_workbook

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

wb = load_workbook(os.path.join(BASE_DIR, "Suivi_collection_cartes.xlsx"))
ws = wb["Collection"]

cards = []
for r in range(2, 200):
    num = ws.cell(r, 2).value
    if not num:
        break
    cards.append({
        "n": num,
        "nom": ws.cell(r, 3).value,
        "rarete": ws.cell(r, 5).value or "",
        "type": ws.cell(r, 6).value or "",
        "essence": ws.cell(r, 7).value or "",
        "ext": ws.cell(r, 4).value or "Tout Commence Ici",
        "note": ws.cell(r, 17).value or "",
    })

print(len(cards), "cartes")
DATA = json.dumps(cards, ensure_ascii=False, separators=(",", ":"))

HTML = r"""<!DOCTYPE html>
<html lang="fr">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1, viewport-fit=cover">
<meta name="apple-mobile-web-app-capable" content="yes">
<meta name="theme-color" content="#151020">
<link rel="apple-touch-icon" href="logo/icone.png">
<link rel="icon" href="logo/icone.png">
<title>Ma collection — Tout Commence Ici</title>
<style>
:root{
  --bg:#151020; --bg-2:#1D1630; --surface:#241C3A; --surface-2:#2E2448;
  --line:#3A2F58; --ink:#F3EEFA; --muted:#A99CC4; --dim:#7C6E9B;
  --rose:#FF4D9D; --orange:#FF8A3D; --jade:#2FD6A6; --violette:#7C4DFF;
  --jaune:#FFD23F; --mauve:#D08BE8; --verte:#58C74E; --bleue:#3B82F6;
  --azur:#22D3EE; --neutre:#9AA4B2;
  --or:#E9C46A; --ok:#4ADE80; --warn:#FBAE3C; --none:#5B4F7A;
  --r:10px;
  --display:"Bahnschrift","Avenir Next Condensed","Roboto Condensed","Segoe UI",system-ui,sans-serif;
  --body:-apple-system,BlinkMacSystemFont,"Segoe UI Variable Text","Segoe UI",Roboto,system-ui,sans-serif;
  --shadow-sm:0 1px 2px rgba(6,3,14,.5);
  --shadow-md:0 6px 18px -4px rgba(6,3,14,.55), 0 2px 6px rgba(6,3,14,.4);
  --shadow-lg:0 18px 40px -12px rgba(6,3,14,.65), 0 4px 12px rgba(6,3,14,.4);
  --shadow-inset:inset 0 1px 0 rgba(255,255,255,.05), inset 0 -1px 0 rgba(0,0,0,.2);
}
*{box-sizing:border-box}
html,body{margin:0;padding:0}
body{
  background:
    radial-gradient(1100px 620px at 18% -8%, rgba(124,77,255,.16), transparent 60%),
    radial-gradient(900px 560px at 100% 0%, rgba(255,77,157,.10), transparent 55%),
    var(--bg);
  background-attachment:fixed;
  color:var(--ink); font-family:var(--body);
  font-size:15px; line-height:1.5; -webkit-text-size-adjust:100%;
  padding-bottom:env(safe-area-inset-bottom);
}
h1,h2,h3{font-family:var(--display); font-weight:600; letter-spacing:.01em; margin:0}
button,input,select,textarea{font:inherit; color:inherit}
:focus-visible{outline:2px solid var(--rose); outline-offset:2px; border-radius:4px}
@media (prefers-reduced-motion:reduce){*{transition:none!important; animation:none!important}}

/* ---------- accueil : choix de l'extension ---------- */
#home{
  min-height:100vh; min-height:100dvh;
  padding:32px 18px calc(32px + env(safe-area-inset-bottom));
  background:radial-gradient(1100px 620px at 50% -10%, rgba(124,77,255,.18), transparent 60%), var(--bg);
}
#home:not([hidden]){display:flex; align-items:center; justify-content:center}
.home-wrap{max-width:640px; width:100%; text-align:center; display:grid; gap:8px; justify-items:center}
#homeBrand{height:120px; width:auto; max-width:80vw; object-fit:contain;
  filter:drop-shadow(0 10px 26px rgba(6,3,14,.55))}
.home-sub{color:var(--muted); margin:0 0 10px; font-size:14.5px}
.ext-tiles{display:grid; gap:12px; width:100%; grid-template-columns:repeat(auto-fit,minmax(220px,1fr))}
.ext-tile{
  background:linear-gradient(160deg, var(--surface-2), var(--surface) 140%); border:1px solid var(--line);
  border-radius:14px; padding:20px 18px; cursor:pointer; text-align:left; color:inherit;
  box-shadow:var(--shadow-md), var(--shadow-inset); transition:transform .16s ease, box-shadow .16s ease, border-color .15s;
  display:grid; gap:10px;
}
.ext-tile:hover{transform:translateY(-3px); box-shadow:var(--shadow-lg), var(--shadow-inset); border-color:var(--dim)}
.ext-tile-logo{height:64px; width:auto; max-width:100%; object-fit:contain; justify-self:center}
.ext-tile-logo-txt{font-size:17px; font-weight:600}
.ext-tile .track{height:8px; background:rgba(6,3,14,.4); border-radius:4px; overflow:hidden; box-shadow:var(--shadow-inset)}
.ext-tile .track i{display:block; height:100%; border-radius:4px; background:var(--rose); box-shadow:0 0 8px -1px var(--rose)}
.ext-tile .ext-pct{color:var(--muted); font-size:13px; font-variant-numeric:tabular-nums}
.ext-info{color:var(--muted); font-size:12.5px; margin:0; line-height:1.4}
.logo-preview{max-height:110px; min-height:60px; display:grid; place-items:center; background:var(--surface);
  border:1px solid var(--line); border-radius:var(--r); padding:8px; overflow:hidden}
.logo-preview img{max-width:100%; max-height:94px; object-fit:contain}
.logo-preview span{color:var(--dim); font-size:12.5px}

/* ---------- en-tête ---------- */
header{
  position:sticky; top:0; z-index:30; background:rgba(21,16,32,.94);
  backdrop-filter:saturate(1.4) blur(10px); border-bottom:1px solid var(--line);
  padding:8px 18px; padding-top:calc(8px + env(safe-area-inset-top));
  box-shadow:var(--shadow-md);
}
header::before{
  content:""; position:absolute; inset:0; pointer-events:none; z-index:-1; border-radius:inherit;
  background:radial-gradient(560px 200px at 50% -40%, rgba(255,77,157,.16), transparent 70%);
}
header::after{
  content:""; position:absolute; left:0; right:0; bottom:0; height:1px;
  background:linear-gradient(to right, transparent, var(--rose) 20%, var(--violette) 55%, var(--azur) 85%, transparent);
  opacity:.55;
}
header.banner{
  background-image:linear-gradient(to right, rgba(21,16,32,.97) 0%, rgba(21,16,32,.72) 55%, rgba(21,16,32,.9) 100%), var(--banner);
  background-size:cover; background-position:center;
}
#brand{height:76px; width:auto; max-width:min(60vw,900px); object-fit:contain; display:block;
  filter:drop-shadow(0 4px 10px rgba(6,3,14,.5));
}
@media (max-width:640px){#brand{height:58px}}
.masthead{display:flex; flex-direction:row; align-items:center; gap:14px;
  max-width:1400px; margin:0 auto; text-align:left}
h1{font-size:19px}
h1 span{color:var(--dim); font-weight:400}
.count{font-family:var(--display); font-size:15px; color:var(--muted);
  font-variant-numeric:tabular-nums; margin-left:auto}
.count b{
  color:var(--ink); font-size:19px; font-weight:700;
  background:linear-gradient(135deg, #fff, var(--rose) 140%);
  -webkit-background-clip:text; background-clip:text; -webkit-text-fill-color:transparent;
}

/* bande des 127 cartes */
.strip{
  display:flex; gap:2px; max-width:1400px; margin:12px auto 0; height:16px;
  padding:3px; background:rgba(6,3,14,.35); border-radius:8px; box-shadow:var(--shadow-inset);
}
.strip i{
  flex:1 1 0; border-radius:2px; background:var(--none); opacity:.5;
  transition:opacity .15s, transform .15s, box-shadow .15s; cursor:pointer; min-width:0;
}
.strip i.has{opacity:1; box-shadow:0 0 6px -1px currentColor}
.strip i:hover{transform:scaleY(1.45)}

/* ---------- barre d'outils ---------- */
.tools{max-width:1400px; margin:0 auto; padding:14px 18px; display:grid; gap:12px}
.toolrow{display:flex; gap:10px; align-items:center; flex-wrap:wrap}
.toolrow .spacer{flex:1 1 0}
.field{position:relative; flex:1 1 auto; min-width:0}
.field input{
  width:100%; background:var(--surface); border:1px solid var(--line);
  border-radius:999px; padding:11px 16px 11px 42px; box-shadow:var(--shadow-inset);
  transition:border-color .15s, box-shadow .15s;
}
.field input:focus-visible{border-color:var(--rose)}
.field input::placeholder{color:var(--dim)}
.field svg{position:absolute; left:15px; top:50%; transform:translateY(-50%); pointer-events:none}
.ic{width:17px; height:17px; flex:none; fill:none; stroke:currentColor; stroke-width:2;
  stroke-linecap:round; stroke-linejoin:round}
.field .ic{stroke:var(--dim); width:18px; height:18px}
.btn.pill{border-radius:999px; padding:10px 18px}
.btn.withic{display:inline-flex; align-items:center; gap:7px}
.btn.primary{background:linear-gradient(160deg, var(--rose), #d93d86 140%); color:#18101C;
  border-color:transparent; font-weight:600}
.btn.primary:hover{color:#18101C; border-color:transparent; filter:brightness(1.07)}
.seg button.icon{padding:9px 12px; display:inline-flex; align-items:center}
.sizectl{display:inline-flex; align-items:center; gap:8px; color:var(--dim)}
.sizectl .ic{width:15px; height:15px}
.sizectl .ic.big{width:20px; height:20px}
.sizectl input[type=range]{width:110px; accent-color:var(--rose); cursor:pointer}
@media (max-width:640px){
  .toolrow .spacer{display:none}
  .toolrow .btn.withic{flex:1 1 0; justify-content:center}
  .toolrow .field{flex:1 1 0}
  .toolrow .btn#toggleF{flex:0 0 auto}
  .toolrow.second{display:grid; grid-template-columns:1fr auto; gap:10px}
  .toolrow.second #sort{min-width:0}
  .toolrow.second #quickMode{order:3}
  .toolrow.second #addCard{order:4}
  .toolrow.second .sizectl{order:5; grid-column:1 / -1}
  .sizectl input[type=range]{flex:1 1 auto; width:auto}
}
select{background:var(--surface); border:1px solid var(--line); border-radius:var(--r); padding:9px 10px;
  box-shadow:var(--shadow-inset)}
.seg{display:flex; border:1px solid var(--line); border-radius:var(--r); overflow:hidden; box-shadow:var(--shadow-sm)}
.seg button{background:var(--surface); border:0; padding:9px 13px; cursor:pointer; color:var(--muted); transition:background .15s, color .15s}
.seg button[aria-pressed=true]{background:linear-gradient(155deg, var(--surface-2), var(--surface)); color:var(--ink); box-shadow:var(--shadow-inset)}

.filters{max-width:1400px; margin:0 auto; padding:2px 18px 16px; display:grid; gap:14px}
.filters[hidden]{display:none}
.fgroup h4{font-size:12px; color:var(--dim); font-weight:500; margin:0 0 7px; font-family:var(--body)}
.filters .chips{max-width:none; margin:0; padding:0}
#fCount{margin-left:7px; background:var(--rose); color:#18101C; border-radius:999px;
  padding:1px 7px; font-size:12px; font-weight:600; display:none}
#fCount.on{display:inline}
.chips{max-width:1400px; margin:0 auto; padding:0 18px 14px; display:flex; gap:7px; flex-wrap:wrap}
.chip{
  background:var(--surface); border:1px solid var(--line); color:var(--muted);
  border-radius:999px; padding:5px 12px; cursor:pointer; font-size:13px;
  display:inline-flex; align-items:center; gap:6px; transition:all .15s;
}
.chip[aria-pressed=true]{background:var(--surface-2); color:var(--ink); border-color:var(--dim); box-shadow:var(--shadow-sm)}
.dot{width:9px; height:9px; border-radius:50%; background:currentColor; flex:none; box-shadow:0 0 5px -1px currentColor}

/* ---------- tableau de bord ---------- */
.tabs{max-width:1400px; margin:0 auto; padding:8px 18px 0}
.tabnav{display:flex; gap:2px; border-bottom:1px solid var(--line)}
@media (max-width:480px){.tabnav button{padding:12px 11px; font-size:14px}}
.tabnav button{background:none; border:0; padding:12px 18px; cursor:pointer; color:var(--dim);
  font-size:14.5px; position:relative; white-space:nowrap; transition:color .15s}
.tabnav button:hover{color:var(--ink)}
.tabnav button[aria-pressed=true]{color:var(--ink); font-weight:600}
.tabnav button[aria-pressed=true]::after{content:""; position:absolute; left:12px; right:12px; bottom:-1px;
  height:2px; background:var(--rose); border-radius:2px; box-shadow:0 0 8px var(--rose)}
.tabnav button[hidden]{display:none}
.xchips{display:flex; gap:8px; flex-wrap:wrap; align-items:center}
.xchips .chip{padding:7px 16px; font-size:14px}
.xchips .chip[aria-pressed=true]{background:var(--rose); color:#18101C; border-color:transparent; font-weight:600}
.xfilters{display:flex; gap:7px; flex-wrap:wrap}
.xlists{display:grid; gap:12px; grid-template-columns:repeat(auto-fit,minmax(280px,1fr))}
.xbox{background:linear-gradient(160deg, var(--bg-2), var(--surface) 130%); border:1px solid var(--line);
  border-radius:var(--r); padding:13px 15px; box-shadow:var(--shadow-md), var(--shadow-inset)}
.xhead{display:flex; justify-content:space-between; align-items:center; gap:10px; margin-bottom:6px}
.xhead b{font-weight:600}
.btn.small{padding:5px 12px; font-size:13px}
.xrow{display:flex; justify-content:space-between; gap:12px; padding:7px 0; border-top:1px solid var(--line); font-size:13.5px}
.xrow .n{color:var(--dim); font-family:var(--display); margin-right:9px; font-variant-numeric:tabular-nums}
.xrow em{font-style:normal; color:var(--muted); font-size:12.5px; text-align:right}
.xempty{color:var(--dim); margin:6px 0 0; font-size:13px}
.xmore{background:none; border:0; color:var(--rose); cursor:pointer; padding:8px 0 0; font-size:13px}
.xtable{min-width:0}
.xtable td{white-space:normal}
.st{display:inline-block; border-radius:6px; padding:2px 9px; font-size:12px; white-space:nowrap}
.st.ok{background:rgba(90,200,79,.16); color:var(--ok)}
.st.part{background:rgba(245,185,74,.16); color:var(--warn)}
.st.no{background:rgba(255,255,255,.06); color:var(--dim)}
.stdup{color:var(--muted); font-size:12px; margin-left:8px; white-space:nowrap}
.cmpwrap{display:grid; gap:14px; grid-template-columns:minmax(0,1fr)}
@media (max-width:480px){
  .stdup{display:block; margin:3px 0 0}
  .xbox{padding:12px 10px}
  .xtable th,.xtable td{padding:8px 5px}
  .xtable{font-size:13px}
  .st{padding:2px 6px}
}
.cmpactions{display:flex; gap:9px; flex-wrap:wrap}
.board{max-width:1400px; margin:18px auto 6px; padding:0 18px; display:grid; gap:12px}
.rings{display:grid; gap:12px; grid-template-columns:repeat(auto-fit,minmax(260px,1fr))}
.minis{display:grid; gap:12px; grid-template-columns:repeat(3,1fr)}
.stat{background:linear-gradient(160deg, var(--bg-2), var(--surface) 130%); border:1px solid var(--line);
  border-radius:var(--r); padding:12px 14px; box-shadow:var(--shadow-md), var(--shadow-inset);
  transition:transform .15s, box-shadow .15s;}
.stat:hover{transform:translateY(-2px); box-shadow:var(--shadow-lg), var(--shadow-inset)}
.stat b{display:block; font-family:var(--display); font-size:22px; font-variant-numeric:tabular-nums; line-height:1.15}
.stat span{color:var(--muted); font-size:12.5px}
.ringcard{display:flex; align-items:center; gap:16px; padding:16px}
.ringcard svg{flex:none; width:96px; height:96px}
.ringcard .rtrack{fill:none; stroke:rgba(255,255,255,.09); stroke-width:10}
.ringcard .rval{fill:none; stroke-width:10; stroke-linecap:round; transition:stroke-dasharray .5s ease;
  filter:drop-shadow(0 0 5px currentColor)}
.ringcard text{fill:var(--ink); font-family:var(--display); font-size:21px; font-weight:600; text-anchor:middle}
.ringcard .rtxt span{display:block}
.ringcard .rtxt .rlab{color:var(--muted); font-size:12.5px}
.ringcard .rtxt .rbig{font-family:var(--display); font-size:26px; font-variant-numeric:tabular-nums; line-height:1.2; color:var(--ink)}

.bars{max-width:1400px; margin:14px auto 0; padding:0 18px; display:grid; gap:14px;
  grid-template-columns:repeat(auto-fit,minmax(290px,1fr))}
.bars section{background:linear-gradient(160deg, var(--bg-2), var(--surface) 130%); border:1px solid var(--line);
  border-radius:var(--r); padding:13px 15px; box-shadow:var(--shadow-md), var(--shadow-inset)}
.bars h3{font-size:13.5px; color:var(--muted); font-weight:500; margin-bottom:9px}
.bars .money{grid-column:1 / -1; display:flex; flex-wrap:wrap; gap:8px 26px; align-items:baseline}
.bars .money h3{flex:1 1 100%; margin:0}
.bars .money b{font-family:var(--display); font-size:19px; font-weight:600; margin-right:6px; font-variant-numeric:tabular-nums}
.bars .money span{color:var(--muted); font-size:13px}
.bars .money b.gain{color:var(--ok)} .bars .money b.loss{color:var(--warn)}
.row{display:grid; grid-template-columns:92px 1fr 58px; align-items:center; gap:9px; margin:5px 0; font-size:13px}
.track{height:7px; background:rgba(6,3,14,.4); border-radius:4px; overflow:hidden; box-shadow:var(--shadow-inset)}
.track i{display:block; height:100%; border-radius:4px; background:linear-gradient(to right, currentColor, currentColor);
  box-shadow:0 0 8px -1px currentColor}
.row em{font-style:normal; color:var(--muted); font-variant-numeric:tabular-nums; text-align:right}
@media (max-width:480px){.minis .stat{padding:10px 11px} .minis .stat b{font-size:19px} .minis .stat span{font-size:11.5px}}

/* ---------- grille ---------- */
main{max-width:1400px; margin:0 auto; padding:16px 18px 60px}
.grid{display:grid; gap:13px; grid-template-columns:repeat(auto-fill,minmax(var(--cardmin,230px),1fr))}
.card{
  background:linear-gradient(160deg, var(--surface), var(--bg-2) 140%);
  border:1px solid var(--line); border-left:3px solid var(--essence,var(--neutre));
  border-radius:var(--r); overflow:hidden; cursor:pointer; text-align:left; padding:0;
  display:flex; flex-direction:column; box-shadow:var(--shadow-md);
  transition:border-color .15s, transform .16s ease, box-shadow .16s ease;
}
.card:hover{
  border-color:var(--dim); border-left-color:var(--essence,var(--neutre));
  transform:translateY(-3px); box-shadow:var(--shadow-lg), 0 0 0 1px var(--essence,transparent);
}
.card.owned{background:linear-gradient(160deg, var(--surface-2), var(--surface) 140%)}
.thumb{aspect-ratio:5/7; background:var(--bg-2); position:relative; display:grid; place-items:center; overflow:hidden}
.thumb::after{
  content:""; position:absolute; inset:0; pointer-events:none;
  background:linear-gradient(to bottom, rgba(0,0,0,0) 65%, rgba(6,3,14,.5) 100%);
}
.thumb img{width:100%; height:100%; object-fit:cover; display:block}
.thumb .ph{font-family:var(--display); font-size:26px; color:var(--dim); font-variant-numeric:tabular-nums}
.thumb img.back{object-fit:cover; opacity:.5}
.card:not(.owned) .thumb img{filter:grayscale(1) brightness(.55)}
.badge{position:absolute; top:7px; right:7px; width:11px; height:11px; border-radius:50%; z-index:1;
  border:2px solid var(--surface); background:var(--none); box-shadow:0 1px 3px rgba(6,3,14,.5)}
.badge.full{background:var(--ok); box-shadow:0 0 7px 0 var(--ok), 0 1px 3px rgba(6,3,14,.5)}
.badge.part{background:var(--warn); box-shadow:0 0 7px 0 var(--warn), 0 1px 3px rgba(6,3,14,.5)}
.meta{padding:8px 10px 10px; display:flex; flex-direction:column; gap:3px; min-height:0; position:relative}
.ess-ico{position:absolute; top:4px; right:7px; width:58px; height:58px; object-fit:contain}
.meta:has(.ess-ico) .num,.meta:has(.ess-ico) .nom,
.meta:has(.ess-ico) .rar{padding-right:62px}
.meta .num{font-family:var(--display); font-size:11.5px; color:var(--dim); font-variant-numeric:tabular-nums}
.meta .nom{font-size:13.5px; line-height:1.25; font-weight:500;
  display:-webkit-box; -webkit-line-clamp:2; -webkit-box-orient:vertical; overflow:hidden}
.meta .rar{font-size:11.5px; color:var(--muted)}
.qty{display:flex; flex-direction:column; gap:5px; padding:0 10px 10px}
.own{display:flex; align-items:center; gap:8px; font-size:12.5px; color:var(--dim); cursor:pointer}
.own input{width:16px; height:16px; accent-color:var(--rose); cursor:pointer; flex:none}
.own.on{color:var(--ink)}
.qrow{display:grid; grid-template-columns:1fr auto auto auto; align-items:center; gap:5px; padding-left:24px}
.qrow .lab{font-size:11px; color:var(--dim); overflow:hidden; text-overflow:ellipsis; white-space:nowrap}
.qrow button{
  width:27px; height:27px; padding:0; border-radius:7px; background:var(--bg-2);
  border:1px solid var(--line); color:var(--muted); cursor:pointer; font-size:15px; line-height:1;
  box-shadow:var(--shadow-sm); transition:color .15s, border-color .15s, transform .1s;
}
.qrow button:hover:not([disabled]){color:var(--ink); border-color:var(--dim)}
.qrow button:active:not([disabled]){transform:scale(.92)}
.qrow button[disabled]{opacity:.3; cursor:default}
.qrow output{min-width:23px; text-align:center; font-family:var(--display); font-size:15px;
  font-variant-numeric:tabular-nums; color:var(--dim)}
.qrow.on output{color:var(--ink)}

/* ---------- tableau ---------- */
.table-wrap{overflow-x:auto; border:1px solid var(--line); border-radius:var(--r)}
table{border-collapse:collapse; width:100%; font-size:13.5px; min-width:820px}
th,td{padding:8px 11px; text-align:left; border-bottom:1px solid var(--line); white-space:nowrap}
th{background:var(--bg-2); color:var(--muted); font-weight:500; position:sticky; top:0}
tbody tr{cursor:pointer}
tbody tr:hover{background:var(--surface)}
td.n{font-family:var(--display); color:var(--dim); font-variant-numeric:tabular-nums}
td.num{font-variant-numeric:tabular-nums}

.empty{padding:56px 18px; text-align:center; color:var(--muted)}
.empty b{display:block; font-family:var(--display); font-size:19px; color:var(--ink); margin-bottom:6px}

/* ---------- panneau détail ---------- */
.scrim{position:fixed; inset:0; background:rgba(10,7,18,.62); z-index:40; display:none}
.scrim.open{display:block}
.panel{
  position:fixed; z-index:50; background:linear-gradient(160deg, var(--bg-2), var(--bg) 140%);
  border-left:1px solid var(--line);
  top:0; right:0; bottom:0; width:min(430px,100%); overflow-y:auto; display:none;
  padding:20px 20px calc(28px + env(safe-area-inset-bottom));
  box-shadow:var(--shadow-lg);
}
.panel.open{display:block}
@media (max-width:640px){
  .panel{top:auto; width:100%; max-height:88vh; border-left:0; border-top:1px solid var(--line);
    border-radius:16px 16px 0 0}
}
.panel .close{position:absolute; top:14px; right:14px; background:var(--surface); border:1px solid var(--line);
  border-radius:8px; width:32px; height:32px; cursor:pointer; color:var(--muted); font-size:17px; line-height:1}
.panel .pnum{font-family:var(--display); color:var(--dim); font-variant-numeric:tabular-nums; font-size:13px}
.pnum-ext-logo{height:22px; width:auto; vertical-align:-6px; object-fit:contain}
.pnum-ext-logo-txt{font-family:var(--display); color:var(--dim); font-size:13px}
.panel h2{font-size:22px; margin:2px 40px 10px 0; line-height:1.2}
.tags{display:flex; gap:7px; flex-wrap:wrap; margin-bottom:16px}
.tag{font-size:12px; padding:3px 9px; border-radius:999px; border:1px solid var(--line); color:var(--muted)}
.tag.ess{color:var(--essence,var(--neutre)); border-color:currentColor}
select.tag{appearance:none; -webkit-appearance:none; -moz-appearance:none; background:var(--surface);
  font:inherit; cursor:pointer}
select.tag.ess option{background:var(--surface-2); color:var(--ink)}
.preview{aspect-ratio:5/7; max-height:280px; margin:0 auto 16px; background:var(--surface);
  border:1px solid var(--line); border-radius:var(--r); display:grid; place-items:center; overflow:hidden;
  box-shadow:var(--shadow-lg)}
.preview img{width:100%; height:100%; object-fit:contain}
.preview span{color:var(--dim); font-size:13px; padding:0 20px; text-align:center}
.pblock{
  background:linear-gradient(160deg, var(--bg-2), var(--surface) 140%); border:1px solid var(--line);
  border-radius:var(--r); padding:14px 15px; box-shadow:var(--shadow-sm), var(--shadow-inset);
  display:grid; gap:12px;
}
.pblock h3{font-size:11px; color:var(--dim); font-weight:600; text-transform:uppercase;
  letter-spacing:.05em; margin:0; font-family:var(--body)}
.ownblk{margin-bottom:0; padding-top:11px}
#pOwn .ownblk:first-child{padding-top:0}
#pOwn .ownblk + .ownblk{border-top:1px solid var(--line)}
.ownblk .own{font-size:14px}
.dup{display:flex; align-items:center; justify-content:space-between; gap:12px;
  margin:8px 0 0 24px; font-size:12.5px; color:var(--muted)}
.dup .stepper{width:132px; flex:none}
.form{display:grid; gap:13px}
.form label:not(.own){display:grid; gap:4px; font-size:12.5px; color:var(--muted)}
.form input:not([type=checkbox]),.form select,.form textarea{
  background:var(--surface); border:1px solid var(--line); border-radius:8px; padding:8px 10px; width:100%;
  box-shadow:var(--shadow-inset);
}
.form textarea{resize:vertical; min-height:62px}
.pair{display:grid; grid-template-columns:1fr 1fr; gap:11px}
.stepper{display:flex; align-items:center; gap:9px}
.stepper button{width:34px; height:34px; border-radius:8px; background:var(--surface);
  border:1px solid var(--line); cursor:pointer; font-size:17px; line-height:1; box-shadow:var(--shadow-sm)}
.stepper output{flex:1; text-align:center; font-family:var(--display); font-size:19px; font-variant-numeric:tabular-nums}
.hint{color:var(--dim); font-size:12px; margin-top:-4px}

/* ---------- pied ---------- */
.btn{background:linear-gradient(160deg, var(--surface-2), var(--surface) 130%); border:1px solid var(--line);
  border-radius:var(--r); padding:9px 14px; cursor:pointer; color:var(--muted);
  box-shadow:var(--shadow-sm), var(--shadow-inset); transition:color .15s, border-color .15s, transform .1s, box-shadow .15s}
.btn:hover{color:var(--ink); border-color:var(--dim); box-shadow:var(--shadow-md), var(--shadow-inset); transform:translateY(-1px)}
.btn:active{transform:translateY(0)}
.btn.warn:hover{color:var(--warn); border-color:var(--warn)}
.btn[aria-pressed=true]{color:var(--ink); border-color:var(--rose); box-shadow:0 0 0 1px var(--rose), var(--shadow-sm)}
.quickbar{max-width:1400px; margin:0 auto; padding:0 18px 14px; color:var(--rose); font-size:13px}
.quickbar[hidden]{display:none}
main.quick .card{cursor:pointer}
main.quick .card:hover{box-shadow:var(--shadow-lg), 0 0 0 2px var(--rose)}
.gear{width:40px; height:40px; border-radius:50%; flex:none; display:grid; place-items:center; cursor:pointer;
  background:linear-gradient(160deg, var(--surface-2), var(--surface) 130%); border:1px solid var(--line);
  color:var(--muted); box-shadow:var(--shadow-sm), var(--shadow-inset);
  transition:color .15s, border-color .15s, transform .3s}
.gear:hover{color:var(--ink); border-color:var(--dim); transform:rotate(60deg)}
.gear .ic{width:20px; height:20px}
#gearHome{position:absolute; top:calc(16px + env(safe-area-inset-top)); right:16px}
#gearApp{position:absolute; top:calc(8px + env(safe-area-inset-top)); right:max(18px, calc((100% - 1400px) / 2 + 18px))}
#home{position:relative}
.gear.alert::after{content:""; position:absolute; top:1px; right:1px; width:11px; height:11px; border-radius:50%;
  background:var(--rose); box-shadow:0 0 0 2px var(--bg)}
.setbtns{display:grid; gap:8px}
.setbtns .btn{text-align:left}
.toast{position:fixed; left:50%; bottom:26px; transform:translateX(-50%); background:var(--surface-2);
  border:1px solid var(--dim); border-radius:999px; padding:9px 18px; z-index:60; display:none; font-size:13.5px;
  box-shadow:var(--shadow-lg)}
.toast.show{display:block}
</style>
</head>
<body>

<section id="home">
  <div class="home-wrap">
    <img id="homeBrand" alt="Gentle Mates TCG" hidden>
    <h1 id="homeTitle">Ma collection</h1>
    <p class="home-sub">Choisis une extension</p>
    <div id="extTiles" class="ext-tiles"></div>
  </div>
  <button class="gear" id="gearHome" aria-label="Réglages" title="Réglages"><svg class="ic" viewBox="0 0 24 24" aria-hidden="true"><path d="M12.22 2h-.44a2 2 0 0 0-2 2v.18a2 2 0 0 1-1 1.73l-.43.25a2 2 0 0 1-2 0l-.15-.08a2 2 0 0 0-2.73.73l-.22.38a2 2 0 0 0 .73 2.73l.15.1a2 2 0 0 1 1 1.72v.51a2 2 0 0 1-1 1.74l-.15.09a2 2 0 0 0-.73 2.73l.22.38a2 2 0 0 0 2.73.73l.15-.08a2 2 0 0 1 2 0l.43.25a2 2 0 0 1 1 1.73V20a2 2 0 0 0 2 2h.44a2 2 0 0 0 2-2v-.18a2 2 0 0 1 1-1.73l.43-.25a2 2 0 0 1 2 0l.15.08a2 2 0 0 0 2.73-.73l.22-.39a2 2 0 0 0-.73-2.73l-.15-.08a2 2 0 0 1-1-1.74v-.5a2 2 0 0 1 1-1.74l.15-.09a2 2 0 0 0 .73-2.73l-.22-.38a2 2 0 0 0-2.73-.73l-.15.08a2 2 0 0 1-2 0l-.43-.25a2 2 0 0 1-1-1.73V4a2 2 0 0 0-2-2z"/><circle cx="12" cy="12" r="3"/></svg></button>
</section>

<div id="app" hidden>

<header id="head">
  <button class="btn" id="backHome" aria-label="Changer d'extension">← Extensions</button>
  <div class="masthead">
    <img id="brand" alt="Gentle Mates TCG" hidden>
    <h1 id="title">Ma collection</h1>
    <div class="count"><b id="hCount">0</b> / __TOTAL__ cartes · <span id="hPct">0 %</span></div>
  </div>
  <button class="gear" id="gearApp" aria-label="Réglages" title="Réglages"><svg class="ic" viewBox="0 0 24 24" aria-hidden="true"><path d="M12.22 2h-.44a2 2 0 0 0-2 2v.18a2 2 0 0 1-1 1.73l-.43.25a2 2 0 0 1-2 0l-.15-.08a2 2 0 0 0-2.73.73l-.22.38a2 2 0 0 0 .73 2.73l.15.1a2 2 0 0 1 1 1.72v.51a2 2 0 0 1-1 1.74l-.15.09a2 2 0 0 0-.73 2.73l.22.38a2 2 0 0 0 2.73.73l.15-.08a2 2 0 0 1 2 0l.43.25a2 2 0 0 1 1 1.73V20a2 2 0 0 0 2 2h.44a2 2 0 0 0 2-2v-.18a2 2 0 0 1 1-1.73l.43-.25a2 2 0 0 1 2 0l.15.08a2 2 0 0 0 2.73-.73l.22-.39a2 2 0 0 0-.73-2.73l-.15-.08a2 2 0 0 1-1-1.74v-.5a2 2 0 0 1 1-1.74l.15-.09a2 2 0 0 0 .73-2.73l-.22-.38a2 2 0 0 0-2.73-.73l-.15.08a2 2 0 0 1-2 0l-.43-.25a2 2 0 0 1-1-1.73V4a2 2 0 0 0-2-2z"/><circle cx="12" cy="12" r="3"/></svg></button>
  <div class="strip" id="strip" role="img" aria-label="Progression carte par carte"></div>
</header>

<div class="tabs">
  <nav class="tabnav" aria-label="Sections">
    <button id="tColl" aria-pressed="true">Collection</button>
    <button id="tStats" aria-pressed="false">Statistiques</button>
    <button id="tCompare" aria-pressed="false">Comparer</button>
    <button id="tAdmin" aria-pressed="false" hidden>Administration</button>
  </nav>
  <input type="file" id="friendFile" accept="application/json" multiple hidden>
</div>

<div class="board" id="board"></div>
<div class="bars" id="bars"></div>

<div class="tools">
  <div class="toolrow">
    <div class="field">
      <svg class="ic" viewBox="0 0 24 24" aria-hidden="true"><circle cx="11" cy="11" r="7"/><path d="M20 20l-3.5-3.5"/></svg>
      <input id="q" type="search" placeholder="Chercher une carte" autocomplete="off" aria-label="Chercher un nom ou un numéro">
    </div>
    <button class="btn pill withic" id="toggleF" aria-expanded="false">
      <svg class="ic" viewBox="0 0 24 24" aria-hidden="true"><path d="M4 6h9M19 6h1M4 12h3M13 12h7M4 18h11M21 18h-1"/><circle cx="16" cy="6" r="2.2"/><circle cx="10" cy="12" r="2.2"/><circle cx="18" cy="18" r="2.2"/></svg>
      Filtres<span id="fCount"></span></button>
  </div>
  <div class="toolrow second">
    <select id="sort" aria-label="Trier">
      <option value="n">Par numéro</option>
      <option value="nom">Par nom</option>
      <option value="rarete">Par rareté</option>
      <option value="val">Par valeur estimée</option>
    </select>
    <div class="seg">
      <button class="icon" id="vGrid" aria-pressed="true" aria-label="Galerie" title="Galerie">
        <svg class="ic" viewBox="0 0 24 24" aria-hidden="true"><rect x="4" y="4" width="6" height="6" rx="1"/><rect x="14" y="4" width="6" height="6" rx="1"/><rect x="4" y="14" width="6" height="6" rx="1"/><rect x="14" y="14" width="6" height="6" rx="1"/></svg></button>
      <button class="icon" id="vList" aria-pressed="false" aria-label="Tableau" title="Tableau">
        <svg class="ic" viewBox="0 0 24 24" aria-hidden="true"><path d="M4 6h16M4 12h16M4 18h16"/></svg></button>
    </div>
    <label class="sizectl" title="Taille des cartes">
      <svg class="ic" viewBox="0 0 24 24" aria-hidden="true"><rect x="7" y="4" width="10" height="16" rx="1.5"/></svg>
      <input id="sizeRange" type="range" min="150" max="340" step="10" value="230" aria-label="Taille des cartes">
      <svg class="ic big" viewBox="0 0 24 24" aria-hidden="true"><rect x="4" y="2" width="16" height="20" rx="2"/></svg>
    </label>
    <span class="spacer"></span>
    <button class="btn withic" id="quickMode" aria-pressed="false">
      <svg class="ic" viewBox="0 0 24 24" aria-hidden="true"><path d="M13 3L5 14h6l-1 7 8-11h-6z"/></svg>
      Mode rapide</button>
    <button class="btn primary withic" id="addCard">
      <svg class="ic" viewBox="0 0 24 24" aria-hidden="true"><path d="M12 5v14M5 12h14"/></svg>
      Ajouter</button>
  </div>
</div>
<div class="quickbar" id="quickbar" hidden>
  Mode rapide actif — touche une carte pour la cocher (encore pour la décocher).
</div>

<div class="filters" id="filters" hidden>
  <div class="fgroup"><h4>Statut</h4><div class="chips" id="fStatus"></div></div>
  <div class="fgroup"><h4>Rareté</h4><div class="chips" id="fRarete"></div></div>
  <div class="fgroup"><h4>Essence</h4><div class="chips" id="fEssence"></div></div>
  <div><button class="btn" id="clearF">Retirer tous les filtres</button></div>
</div>

<main id="out"></main>



</div>

<input type="file" id="file" accept="application/json" hidden>
<div class="scrim" id="scrim"></div>
<aside class="panel" id="panel" aria-label="Détail de la carte"></aside>
<div class="toast" id="toast"></div>

<script>
const IMG = __IMAGES__;
function asset(k){ return IMG[k] || null; }
const BASE = __DATA__;
const KEY = "m8tcg.v1";
const XKEY = "m8tcg.extra";
const OVKEY = "m8tcg.overrides";
const DELKEY = "m8tcg.deleted";
const EXTXKEY = "m8tcg.extsExtra";
const EXTMETAKEY = "m8tcg.extMeta";
let EXTRA = [], OVERRIDES = {}, DELETED = [], EXT_EXTRA = [], EXT_META = {};
try{ EXTRA = JSON.parse(localStorage.getItem(XKEY) || "[]"); }catch(e){ EXTRA = []; }
try{ OVERRIDES = JSON.parse(localStorage.getItem(OVKEY) || "{}"); }catch(e){ OVERRIDES = {}; }
try{ DELETED = JSON.parse(localStorage.getItem(DELKEY) || "[]"); }catch(e){ DELETED = []; }
try{ EXT_EXTRA = JSON.parse(localStorage.getItem(EXTXKEY) || "[]"); }catch(e){ EXT_EXTRA = []; }
try{ EXT_META = JSON.parse(localStorage.getItem(EXTMETAKEY) || "{}"); }catch(e){ EXT_META = {}; }
let CARDS = [];
function rebuildCards(){
  CARDS = BASE.filter(c=>!DELETED.includes(c.n))
    .map(c=> OVERRIDES[c.n] ? Object.assign({}, c, OVERRIDES[c.n]) : c)
    .concat(EXTRA);
}
rebuildCards();
function saveExtra(){
  try{ localStorage.setItem(XKEY, JSON.stringify(EXTRA)); }catch(e){}
  rebuildCards();
  changed();
}
function saveOverrides(){
  try{ localStorage.setItem(OVKEY, JSON.stringify(OVERRIDES)); }catch(e){}
  rebuildCards();
  changed();
}
function saveDeleted(){
  try{ localStorage.setItem(DELKEY, JSON.stringify(DELETED)); }catch(e){}
  rebuildCards();
  changed();
}
function saveExtExtra(){
  try{ localStorage.setItem(EXTXKEY, JSON.stringify(EXT_EXTRA)); }catch(e){}
  changed();
}
function saveExtMeta(){
  try{ localStorage.setItem(EXTMETAKEY, JSON.stringify(EXT_META)); }catch(e){}
  changed();
}
function isExtra(n){ return EXTRA.some(c=>c.n===n); }
const ESS = {Rose:"--rose",Orange:"--orange",Jade:"--jade",Violette:"--violette",Jaune:"--jaune",
  Mauve:"--mauve",Verte:"--verte",Bleue:"--bleue",Azur:"--azur",Neutre:"--neutre"};
const RARS = ["Commune","Peu commune","Rare","Épique","Légendaire","Secrète","Promo","Essence"];
const RARORDER = {}; RARS.forEach((r,i)=>RARORDER[r]=i);
const ESSLIST = Object.keys(ESS);
const TYPES = ["Joueur","Coach / Instant","Objet / Lien","Moment / Instant","Essence"];
const ETATS = ["","Mint","Near Mint","Excellent","Bon","Abîmé"];

/* ---------- stockage ---------- */
let store = {};
let memoryOnly = false;
try{
  const raw = localStorage.getItem(KEY);
  if(raw) store = JSON.parse(raw);
}catch(e){ memoryOnly = true; }
function save(){
  if(!memoryOnly){
    try{ localStorage.setItem(KEY, JSON.stringify(store)); }
    catch(e){ memoryOnly = true; toast("Sauvegarde locale impossible — pense à exporter"); }
  }
  changed();
}
/* ---------- synchronisation (fichier secret « gist » du compte GitHub) ---------- */
const SYNCKEY="m8tcg.sync", DIRTYKEY="m8tcg.syncDirty";
const GFILE="m8tcg-collection.json";
const GDESC="Ma collection Gentle Mates TCG (sauvegarde automatique)";
let SYNC=null, syncDirty=false, dirtySeq=0, syncBusy=false, syncTimer=null, syncConflict=null;
try{ SYNC=JSON.parse(localStorage.getItem(SYNCKEY)||"null"); }catch(e){ SYNC=null; }
try{ syncDirty=localStorage.getItem(DIRTYKEY)==="1"; }catch(e){}
function saveSyncCfg(){
  try{ if(SYNC) localStorage.setItem(SYNCKEY, JSON.stringify(SYNC)); else localStorage.removeItem(SYNCKEY); }catch(e){}
}
function setDirty(v){
  syncDirty=v;
  try{ localStorage.setItem(DIRTYKEY, v?"1":"0"); }catch(e){}
}
let syncText="";
function setSync(t){
  syncText=t; const el=document.getElementById("syncState"); if(el) el.textContent=t;
  const alert=/refusé|introuvable|choix/.test(t);
  document.querySelectorAll(".gear").forEach(g=>g.classList.toggle("alert",alert));
}
function changed(){ dirtySeq++; if(!SYNC) return; setDirty(true); schedulePush(); }
function snapshot(){
  return {v:3, updatedAt:new Date().toISOString(), data:store, extra:EXTRA, overrides:OVERRIDES,
    deleted:DELETED, extsExtra:EXT_EXTRA, extMeta:EXT_META};
}
function ownedCount(d){
  return Object.keys(d||{}).filter(n=>d[n] && (+d[n].qn||0)+(+d[n].qr||0)>0).length;
}
function hasAnyLocal(){
  return ownedCount(store)>0 || EXTRA.length>0 || DELETED.length>0 || EXT_EXTRA.length>0
    || Object.keys(OVERRIDES).length>0;
}
function fmtDate(iso){
  try{ return new Date(iso).toLocaleString("fr-FR",{dateStyle:"medium",timeStyle:"short"}); }catch(e){ return ""; }
}
function refreshAll(){
  rebuildCards(); applyAdminVisibility();
  if(currentExt && !document.getElementById("app").hidden){
    render();
    if(document.getElementById("tCompare").getAttribute("aria-pressed")==="true") renderCompare();
    else if(document.getElementById("tAdmin").getAttribute("aria-pressed")==="true") renderAdmin();
  }
  renderHome();
}
function applySnapshot(j){
  store = (j.data && typeof j.data==="object") ? j.data : {};
  EXTRA = Array.isArray(j.extra) ? j.extra : [];
  OVERRIDES = (j.overrides && typeof j.overrides==="object") ? j.overrides : {};
  DELETED = Array.isArray(j.deleted) ? j.deleted : [];
  EXT_EXTRA = Array.isArray(j.extsExtra) ? j.extsExtra : [];
  EXT_META = (j.extMeta && typeof j.extMeta==="object") ? j.extMeta : {};
  [[KEY,store],[XKEY,EXTRA],[OVKEY,OVERRIDES],[DELKEY,DELETED],[EXTXKEY,EXT_EXTRA],[EXTMETAKEY,EXT_META]]
    .forEach(([k,v])=>{ try{ localStorage.setItem(k, JSON.stringify(v)); }catch(e){} });
  refreshAll();
}
async function gh(token, path, opts){
  opts=opts||{};
  const headers={"Accept":"application/vnd.github+json","Authorization":"Bearer "+token,
    "X-GitHub-Api-Version":"2022-11-28"};
  if(opts.body) headers["Content-Type"]="application/json";
  const r=await fetch("https://api.github.com"+path,
    {method:opts.method||"GET", headers, body:opts.body, cache:"no-store"});
  if(!r.ok){ const err=new Error("HTTP "+r.status); err.status=r.status; throw err; }
  return r.json();
}
async function gistContent(g){
  const f=g.files && g.files[GFILE];
  if(!f){ const err=new Error("fichier absent"); err.status=404; throw err; }
  const text = f.truncated ? await (await fetch(f.raw_url,{cache:"no-store"})).text() : f.content;
  return JSON.parse(text);
}
async function readCloud(){ return gistContent(await gh(SYNC.token,"/gists/"+SYNC.id)); }
async function writeCloud(snap){
  await gh(SYNC.token,"/gists/"+SYNC.id,
    {method:"PATCH", body:JSON.stringify({files:{[GFILE]:{content:JSON.stringify(snap)}}})});
}
function syncError(e){
  if(e && (e.status===401||e.status===403)) setSync("Jeton refusé — remplace-le dans les réglages (roue crantée)");
  else if(e && e.status===404) setSync("Sauvegarde en ligne introuvable — vois les réglages (roue crantée)");
  else setSync("Pas de connexion : tes saisies sont gardées ici, nouvel essai plus tard");
}
function schedulePush(){
  if(!SYNC) return;
  setSync("Modifications à envoyer…");
  clearTimeout(syncTimer); syncTimer=setTimeout(syncNow,1500);
}
async function syncNow(){
  if(!SYNC || syncBusy || syncConflict) return;
  if(navigator.onLine===false){ setSync("Hors ligne : tes saisies sont gardées ici"); return; }
  syncBusy=true; setSync("Synchronisation…");
  try{
    const cloud=await readCloud();
    const cloudChanged = cloud.updatedAt!==SYNC.cloudAt;
    if(cloudChanged && syncDirty){ showConflict(cloud); return; }
    if(cloudChanged){
      applySnapshot(cloud);
      SYNC.cloudAt=cloud.updatedAt; SYNC.last=new Date().toISOString(); saveSyncCfg();
      setSync("Synchronisé"); toast("Collection mise à jour depuis ton compte");
      return;
    }
    if(syncDirty){
      const seq=dirtySeq, snap=snapshot();
      await writeCloud(snap);
      SYNC.cloudAt=snap.updatedAt; saveSyncCfg();
      if(seq===dirtySeq) setDirty(false); else schedulePush();
    }
    SYNC.last=new Date().toISOString(); saveSyncCfg();
    if(!syncDirty) setSync("Synchronisé");
  }catch(e){ syncError(e); }
  finally{ syncBusy=false; }
}
function showConflict(cloud){
  syncConflict=cloud;
  setSync("Un choix est à faire : ouvre les réglages (roue crantée)");
  openSync();
}
async function resolveConflict(keepLocal){
  const cloud=syncConflict; syncConflict=null;
  if(!cloud || !SYNC) return;
  close();
  if(keepLocal){
    SYNC.cloudAt=cloud.updatedAt; saveSyncCfg(); setDirty(true);
    await syncNow();
    toast("Version de cet appareil envoyée");
  }else{
    applySnapshot(cloud);
    SYNC.cloudAt=cloud.updatedAt; SYNC.last=new Date().toISOString(); saveSyncCfg();
    setDirty(false); setSync("Synchronisé"); toast("Version en ligne récupérée");
  }
}
async function connect(token){
  token=(token||"").trim();
  if(!token){ toast("Colle d'abord ton jeton"); return; }
  setSync("Connexion à GitHub…");
  try{
    let found=null;
    for(let page=1; page<=10 && !found; page++){
      const list=await gh(token,"/gists?per_page=100&page="+page);
      found=list.find(g=>g.files && g.files[GFILE]) || null;
      if(list.length<100) break;
    }
    if(found){
      const cloud=await gistContent(await gh(token,"/gists/"+found.id));
      SYNC={token, id:found.id, cloudAt:"", last:""}; saveSyncCfg();
      if(hasAnyLocal()){ showConflict(cloud); return; }
      applySnapshot(cloud);
      SYNC.cloudAt=cloud.updatedAt; SYNC.last=new Date().toISOString(); saveSyncCfg(); setDirty(false);
      setSync("Synchronisé"); toast("Collection récupérée depuis ton compte"); openSync();
    }else{
      const snap=snapshot();
      const g=await gh(token,"/gists",{method:"POST",
        body:JSON.stringify({description:GDESC, public:false, files:{[GFILE]:{content:JSON.stringify(snap)}}})});
      SYNC={token, id:g.id, cloudAt:snap.updatedAt, last:new Date().toISOString()}; saveSyncCfg(); setDirty(false);
      setSync("Synchronisé"); toast("Synchronisation activée"); openSync();
    }
  }catch(e){
    SYNC=null; saveSyncCfg();
    setSync("Mode local : les saisies restent sur cet appareil");
    toast(e && (e.status===401||e.status===403)
      ? "Jeton refusé : vérifie qu'il est copié en entier et qu'il a l'option « gist »"
      : "Impossible de joindre GitHub");
  }
}
function disconnect(){
  SYNC=null; saveSyncCfg(); setDirty(false); syncConflict=null; clearTimeout(syncTimer);
  setSync("Mode local : les saisies restent sur cet appareil");
  close(); toast("Synchronisation désactivée sur cet appareil");
}
function openSync(){
  current=null;
  const p=document.getElementById("panel");
  p.style.setProperty("--essence","var(--neutre)");
  let body;
  if(syncConflict){
    const c=syncConflict;
    body=`<h2>Quelle version garder ?</h2>
      <div class="form">
        <p class="hint" style="margin:0">Cet appareil et ton compte GitHub contiennent chacun des données que l'autre n'a pas.</p>
        <div class="pblock"><h3>Cet appareil</h3><div>${ownedCount(store)} carte(s) possédée(s)</div></div>
        <div class="pblock"><h3>En ligne</h3><div>${ownedCount(c.data)} carte(s) possédée(s)${c.updatedAt?` · modifié le ${fmtDate(c.updatedAt)}`:""}</div></div>
        <button class="btn" id="sKeepLocal">Garder cet appareil</button>
        <button class="btn" id="sKeepCloud">Garder la version en ligne</button>
        <p class="hint">La version que tu ne gardes pas est remplacée. Les anciennes versions en ligne restent consultables dans l'historique du fichier sur GitHub.</p>
      </div>`;
  }else if(SYNC){
    body=`<h2>Synchronisation active</h2>
      <div class="form">
        <p class="hint" style="margin:0">${SYNC.last?`Dernière synchronisation : ${fmtDate(SYNC.last)}`:"Pas encore synchronisé"}</p>
        <p style="margin:0">Ta collection est enregistrée dans un fichier secret de ton compte GitHub, mis à jour tout seul quand tu coches une carte. Sur un autre appareil : ouvre ce site, clique sur la roue crantée puis sur « Synchroniser mes appareils » et colle le même jeton.</p>
        <button class="btn" id="sNow">Synchroniser maintenant</button>
        <a class="btn" style="text-align:center;text-decoration:none" href="https://gist.github.com/${SYNC.id}" target="_blank" rel="noopener">Voir le fichier et son historique</a>
        <button class="btn warn" id="sOff">Désactiver sur cet appareil</button>
        <p class="hint">Désactiver garde tes cartes ici et retire seulement le jeton de cet appareil.</p>
      </div>`;
  }else{
    body=`<h2>Synchroniser mes appareils</h2>
      <div class="form">
        <p style="margin:0">1. <a href="https://github.com/settings/tokens/new?scopes=gist&description=m8-collection" target="_blank" rel="noopener" style="color:var(--rose)">Crée un jeton sur GitHub</a> : l'option « gist » est déjà cochée, ne coche rien d'autre. Choisis une durée (« No expiration » pour ne pas avoir à recommencer), puis « Generate token ».</p>
        <p style="margin:0">2. Copie le jeton (il commence par <code>ghp_</code>) et colle-le ici :</p>
        <label>Jeton<input id="sToken" type="password" autocomplete="off" placeholder="ghp_…"></label>
        <button class="btn" id="sGo">Activer</button>
        <p class="hint">Le jeton reste sur cet appareil. Ne le partage avec personne.</p>
      </div>`;
  }
  p.innerHTML=`<button class="close" id="pClose" aria-label="Fermer">×</button>
    <div class="pnum">Synchronisation</div>${body}`;
  document.getElementById("pClose").onclick=close;
  const on=(id,fn)=>{ const el=document.getElementById(id); if(el) el.onclick=fn; };
  on("sKeepLocal",()=>resolveConflict(true));
  on("sKeepCloud",()=>resolveConflict(false));
  on("sNow",async()=>{ await syncNow(); openSync(); });
  on("sOff",disconnect);
  on("sGo",async()=>{ await connect(document.getElementById("sToken").value); });
  p.classList.add("open"); document.getElementById("scrim").classList.add("open");
}
document.addEventListener("visibilitychange",()=>{ syncNow(); });
window.addEventListener("online",()=>{ syncNow(); });

function S(n){ return store[n] || (store[n] = {qn:0,qr:0,rev:"",etat:"",prix:"",val:"",date:"",loc:"",notes:""}); }
function has(n){ const s=store[n]; return s ? (+s.qn||0)+(+s.qr||0) > 0 : false; }
const NOREV_RARETE=new Set(["Épique","Légendaire","Secrète","Promo"]);
function autoRev(c){ return c ? !NOREV_RARETE.has(c.rarete) : true; }
function effRev(c,s){
  if(s){ if(s.rev==="Oui") return true; if(s.rev==="Non") return false; }
  return autoRev(c);
}
function effEssence(c,s){ return (s && s.essence) || (c && c.essence) || ""; }
function statut(n){ return statutOf(store,n); }
function statutOf(data,n){
  const s = data[n]; if(!s) return "manquante";
  const qn=+s.qn||0, qr=+s.qr||0;
  if(qn+qr===0) return "manquante";
  if(effRev(CARDS.find(c=>c.n===n),s) && (qn===0 || qr===0)) return "partielle";
  return "complete";
}
/* état d'une carte en mots : Complète / Sans reverse / Sans normale / Absente */
function versionState(data,c){
  const s=data[c.n]||{}, n=+s.qn||0, r=+s.qr||0;
  if(n+r===0) return {cls:"no", txt:"Absente"};
  if(!effRev(c,s) || (n>0&&r>0)) return {cls:"ok", txt:"Complète"};
  return n>0 ? {cls:"part", txt:"Sans reverse"} : {cls:"part", txt:"Sans normale"};
}
function dupOf(data,n){
  const s=data[n]; if(!s) return 0;
  return Math.max(0,(+s.qn||0)-1) + Math.max(0,(+s.qr||0)-1);
}

/* ---------- extensions ---------- */
const EXTKEY="m8tcg.ext";
let currentExt=null;
const ADMKEY="m8tcg.admin";
let adminUnlocked=false;
try{ adminUnlocked = localStorage.getItem(ADMKEY)==="1"; }catch(e){}
function applyAdminVisibility(){
  if(!adminUnlocked && document.getElementById("tAdmin").getAttribute("aria-pressed")==="true") setTab("coll");
}
function toggleAdmin(){
  adminUnlocked=!adminUnlocked;
  try{ localStorage.setItem(ADMKEY, adminUnlocked?"1":"0"); }catch(e){}
  applyAdminVisibility();
  toast(adminUnlocked ? "Mode administration activé" : "Mode administration désactivé");
}
function armSecretUnlock(el){
  let taps=0, timer=null;
  el.addEventListener("click",()=>{
    taps++;
    clearTimeout(timer);
    timer=setTimeout(()=>{ taps=0; }, 2000);
    if(taps>=5){ taps=0; clearTimeout(timer); toggleAdmin(); }
  });
}
function extsList(){ return [...new Set([...CARDS.map(c=>c.ext||"Tout Commence Ici"), ...EXT_EXTRA])]; }
function extSlug(ext){
  return String(ext).toLowerCase().normalize("NFD").replace(/[̀-ͯ]/g,"").replace(/[^a-z0-9]/g,"");
}
function extFallback(img){
  img.replaceWith(Object.assign(document.createElement("span"),{className:img.dataset.fallbackClass||"",textContent:img.alt}));
}
function extLogo(ext, cls){
  const meta=EXT_META[ext];
  if(meta && meta.logo) return `<img class="${cls}" src="${meta.logo}" alt="${ext}">`;
  const slug=extSlug(ext);
  const src=asset("ext-"+slug) || `logo/${slug}.png`;
  return `<img class="${cls}" data-fallback-class="${cls}-txt" src="${src}" alt="${ext}" onerror="extFallback(this)">`;
}
function renderHome(){
  const host=document.getElementById("extTiles");
  host.innerHTML = extsList().map(ext=>{
    const set = CARDS.filter(c=>(c.ext||"Tout Commence Ici")===ext);
    const got = set.filter(c=>has(c.n)).length;
    const pct = set.length ? Math.round(got/set.length*100) : 0;
    const info = EXT_META[ext] && EXT_META[ext].info;
    return `<button class="ext-tile" data-ext="${ext}">
      ${extLogo(ext,"ext-tile-logo")}
      ${info?`<p class="ext-info">${info}</p>`:""}
      <div class="track"><i style="width:${pct}%"></i></div>
      <span class="ext-pct">${got} / ${set.length} cartes · ${pct} %</span>
    </button>`;
  }).join("");
  host.querySelectorAll("[data-ext]").forEach(b=>b.onclick=()=>enterExt(b.dataset.ext));
}
function enterExt(ext){
  currentExt=ext;
  try{ localStorage.setItem(EXTKEY, ext); }catch(e){}
  document.getElementById("home").hidden=true;
  document.getElementById("app").hidden=false;
  render();
}
function goHome(){
  document.getElementById("app").hidden=true;
  document.getElementById("home").hidden=false;
  renderHome();
}

/* ---------- filtres ---------- */
let view="grid", sortBy="n", query="", quick=false;
let FRIENDS=[];
let fSt=new Set(), fRa=new Set(), fEs=new Set();

function chip(host, label, active, onclick, color){
  const b=document.createElement("button");
  b.className="chip"; b.setAttribute("aria-pressed", active);
  if(color){ const d=document.createElement("span"); d.className="dot";
    d.style.color=`var(${color})`; b.appendChild(d); }
  b.appendChild(document.createTextNode(label));
  b.onclick=onclick; host.appendChild(b); return b;
}
function buildFilters(){
  const st=document.getElementById("fStatus"); st.innerHTML="";
  [["manquante","Manquantes"],["partielle","Une version manque"],["complete","Complètes"]]
    .forEach(([k,l])=>chip(st,l,fSt.has(k),()=>{fSt.has(k)?fSt.delete(k):fSt.add(k); render();}));
  const ra=document.getElementById("fRarete"); ra.innerHTML="";
  RARS.forEach(r=>chip(ra,r,fRa.has(r),()=>{fRa.has(r)?fRa.delete(r):fRa.add(r); render();}));
  const es=document.getElementById("fEssence"); es.innerHTML="";
  ESSLIST.forEach(e=>chip(es,e,fEs.has(e),()=>{fEs.has(e)?fEs.delete(e):fEs.add(e); render();},ESS[e]));
}

function visible(){
  const q=query.trim().toLowerCase();
  return CARDS.filter(c=>{
    if(currentExt && (c.ext||"Tout Commence Ici")!==currentExt) return false;
    if(q && !(c.nom.toLowerCase().includes(q) || c.n.toLowerCase().includes(q))) return false;
    if(fRa.size && !fRa.has(c.rarete)) return false;
    if(fEs.size && !fEs.has(effEssence(c,store[c.n]))) return false;
    if(fSt.size && !fSt.has(statut(c.n))) return false;
    return true;
  }).sort((a,b)=>{
    if(sortBy==="nom") return a.nom.localeCompare(b.nom,"fr");
    if(sortBy==="rarete"){
      const d=(RARORDER[b.rarete]??-1)-(RARORDER[a.rarete]??-1);
      return d || a.n.localeCompare(b.n);
    }
    if(sortBy==="val") return (+(S(b.n).val)||0)-(+(S(a.n).val)||0) || a.n.localeCompare(b.n);
    return a.n.localeCompare(b.n);
  });
}

/* ---------- images ---------- */
function cardFallback(img, n){
  if(window.CARDBACK && !img.dataset.tried){
    img.dataset.tried="1"; img.classList.add("back"); img.src=window.CARDBACK; return;
  }
  img.replaceWith(Object.assign(document.createElement("span"),{className:"ph",textContent:n}));
}
function cardSrc(n){ return asset(n) || `images-cartes/${n}.png`; }
function thumbHTML(c){
  return `<img src="${cardSrc(c.n)}" alt="" loading="lazy"
    onerror="cardFallback(this,'${c.n}')">`;
}

function essIcon(c){
  const ess=effEssence(c,store[c.n]);
  if(!ess) return "";
  const slug=ess.toLowerCase();
  const isrc=asset("essence-"+slug) || `logo/essence-${slug}.png`;
  return `<img class="ess-ico" src="${isrc}" alt="${ess}"
    title="${ess}" loading="lazy" onerror="this.remove()">`;
}

/* ---------- habillage fourni dans le dossier logo ---------- */
function probe(src){
  return new Promise(res=>{ const i=new Image(); i.onload=()=>res(src); i.onerror=()=>res(null); i.src=src; });
}
async function first(list){
  for(const p of list){ const r=await probe(p); if(r) return r; }
  return null;
}
async function loadBrand(){
  const logo=asset("logo") || await first(["logo/logo.svg","logo/logo.png","logo/logo.webp"]);
  if(logo){
    const b=document.getElementById("brand");
    b.src=logo; b.hidden=false;
    document.getElementById("title").hidden=true;
    const hb=document.getElementById("homeBrand");
    hb.src=logo; hb.hidden=false;
    document.getElementById("homeTitle").hidden=true;
  }
  const ban=asset("banniere") || await first(["logo/banniere.jpg","logo/banniere.png","logo/banniere.webp"]);
  if(ban){
    const h=document.getElementById("head");
    h.style.setProperty("--banner",`url("${ban}")`);
    h.classList.add("banner");
  }
  const back=asset("dos-carte") || await first(["logo/dos-carte.png","logo/dos-carte.jpg","logo/dos-carte.webp"]);
  if(back){ window.CARDBACK=back; render(); }
}

/* ---------- bloc quantités d'une vignette ---------- */
function qtyHTML(c){
  const s=S(c.n);
  const line=(k,lab)=>{
    const v=+s[k]||0, on=v>0, dup=Math.max(0,v-1);
    if(k==="qr" && !effRev(c,s)) return "";
    return `<label class="own${on?" on":""}" data-own="${k}">
        <input type="checkbox"${on?" checked":""}><span>${lab}</span></label>`
      + (on?`<div class="qrow${dup>0?" on":""}" data-row="${k}">
        <span class="lab">Doublons</span>
        <button data-k="${k}" data-d="-1" aria-label="Un doublon de moins"${dup===0?" disabled":""}>−</button>
        <output>${dup}</output>
        <button data-k="${k}" data-d="1" aria-label="Un doublon de plus">+</button></div>`:"");
  };
  return line("qn","Normale") + line("qr","Reverse");
}
function bindQty(el, c){
  el.querySelectorAll(".own").forEach(lab=>{
    lab.onclick=e=>e.stopPropagation();
    lab.querySelector("input").onchange=ev=>{
      S(c.n)[lab.dataset.own] = ev.target.checked ? 1 : 0;
      save();
      if(fSt.size || sortBy==="val"){ render(); return; }
      paintQty(el,c); stats();
    };
  });
  el.querySelectorAll(".qrow button").forEach(b=>{
    b.onclick=e=>{
      e.stopPropagation();
      const k=b.dataset.k;
      S(c.n)[k]=Math.max(1,(+S(c.n)[k]||1)+(+b.dataset.d));
      save(); paintQty(el,c); stats();
    };
  });
}
function toggleQuick(c){
  const s=S(c.n);
  s.qn = (+s.qn||0)>0 ? 0 : 1;
  save(); render();
}
function paintQty(el, c){
  el.querySelector(".qty").innerHTML=qtyHTML(c);
  bindQty(el,c);
  const st=statut(c.n);
  el.classList.toggle("owned", st!=="manquante");
  const badge=el.querySelector(".badge");
  badge.className="badge"+(st==="complete"?" full":st==="partielle"?" part":"");
}

/* ---------- rendu ---------- */
function render(){
  const list=visible(), out=document.getElementById("out");
  if(!list.length){
    out.innerHTML=`<div class="empty"><b>Aucune carte ne correspond</b>Élargis ta recherche ou retire un filtre.</div>`;
  } else if(view==="grid"){
    const g=document.createElement("div"); g.className="grid";
    list.forEach(c=>{
      const st=statut(c.n);
      const el=document.createElement("div");
      el.className="card"+(st!=="manquante"?" owned":"");
      el.style.setProperty("--essence",`var(${ESS[effEssence(c,store[c.n])]||"--neutre"})`);
      el.tabIndex=0; el.setAttribute("role","button");
      el.innerHTML=`
        <div class="thumb">${thumbHTML(c)}
          <span class="badge ${st==="complete"?"full":st==="partielle"?"part":""}"></span></div>
        <div class="meta">
          ${essIcon(c)}
          <span class="num">${c.n}</span>
          <span class="nom">${c.nom}</span>
          <span class="rar">${c.rarete||"Rareté inconnue"}</span>
        </div>
        <div class="qty">${qtyHTML(c)}</div>`;
      bindQty(el,c);
      el.onclick=()=>{ if(quick){ toggleQuick(c); } else { open(c.n); } };
      el.onkeydown=e=>{ if(e.key==="Enter"||e.key===" "){e.preventDefault();
        if(quick){ toggleQuick(c); } else { open(c.n); } } };
      g.appendChild(el);
    });
    out.innerHTML=""; out.appendChild(g);
  } else {
    const rows=list.map(c=>{
      const s=S(c.n);
      const vs=versionState(store,c);
      return `<tr data-n="${c.n}">
        <td class="n">${c.n}</td><td>${c.nom}</td><td>${c.rarete||"—"}</td>
        <td>${c.type||"—"}</td><td style="color:var(${ESS[effEssence(c,s)]||"--neutre"})">${effEssence(c,s)||"—"}</td>
        <td class="num">${+s.qn||0}</td><td class="num">${+s.qr||0}</td>
        <td>${effRev(c,s)?"Oui":"Non"}${s.rev?"":" (auto)"}</td><td>${s.etat||"—"}</td>
        <td class="num">${s.prix!==""?(+s.prix).toFixed(2)+" €":"—"}</td>
        <td class="num">${s.val!==""?(+s.val).toFixed(2)+" €":"—"}</td>
        <td>${s.loc||"—"}</td><td><span class="st ${vs.cls}">${vs.txt}</span></td></tr>`;
    }).join("");
    out.innerHTML=`<div class="table-wrap"><table>
      <thead><tr><th>N°</th><th>Nom</th><th>Rareté</th><th>Type</th><th>Essence</th>
      <th>Norm.</th><th>Rev.</th><th>Reverse ?</th><th>État</th><th>Prix payé</th>
      <th>Valeur</th><th>Emplacement</th><th>Statut</th></tr></thead><tbody>${rows}</tbody></table></div>`;
    out.querySelectorAll("tbody tr").forEach(tr=>tr.onclick=()=>open(tr.dataset.n));
  }
  stats();
  updateFCount();
}

/* ---------- statistiques ---------- */
function stats(){
  const owned=CARDS.filter(c=>has(c.n)).length;
  const total=CARDS.length;
  let qn=0,qr=0,paid=0,val=0;
  CARDS.forEach(c=>{ const s=store[c.n]; if(!s) return;
    qn+=+s.qn||0; qr+=+s.qr||0; paid+=+s.prix||0; val+=+s.val||0;
  });
  const pct=total?Math.round(owned/total*100):0;
  const revElig=CARDS.filter(c=>effRev(c,store[c.n]));
  const revHave=revElig.filter(c=>(+((store[c.n]||{}).qr)||0)>0).length;
  // Master set : chaque version (normale + reverse quand elle existe) compte comme une pièce à part.
  const normalHave=CARDS.filter(c=>(+((store[c.n]||{}).qn)||0)>0).length;
  const masterTotal=total+revElig.length;
  const masterHave=normalHave+revHave;
  const pctMaster=masterTotal?Math.round(masterHave/masterTotal*100):0;
  document.getElementById("hCount").textContent=owned;
  document.getElementById("hPct").textContent=pct+" %";

  const eur=n=>n.toLocaleString("fr-FR",{minimumFractionDigits:2,maximumFractionDigits:2})+" €";
  const diff=val-paid;
  let dups=0; CARDS.forEach(c=>{ dups+=dupOf(store,c.n); });
  const ring=(p,color,label,big,sub)=>{
    const C=2*Math.PI*42;
    return `<div class="stat ringcard">
      <svg viewBox="0 0 100 100" role="img" aria-label="${label} : ${p} %">
        <circle class="rtrack" cx="50" cy="50" r="42"/>
        <circle class="rval" cx="50" cy="50" r="42" style="stroke:${color}; color:${color}"
          stroke-dasharray="${(C*p/100).toFixed(1)} ${C.toFixed(1)}" transform="rotate(-90 50 50)"/>
        <text x="50" y="57">${p} %</text>
      </svg>
      <div class="rtxt"><span class="rlab">${label}</span><span class="rbig">${big}</span><span class="rlab">${sub}</span></div>
    </div>`;
  };
  document.getElementById("board").innerHTML=`
    <div class="rings">
      ${ring(pct,"var(--rose)","Full set",`${owned} / ${total}`,"cartes différentes")}
      ${ring(pctMaster,"var(--violette)","Master set",`${masterHave} / ${masterTotal}`,"normales + reverses")}
    </div>
    <div class="minis">
      <div class="stat"><b>${qn + qr}</b><span>Cartes ouvertes</span></div>
      <div class="stat"><b>${revHave} / ${revElig.length}</b><span>Reverses trouvées</span></div>
      <div class="stat"><b>${dups}</b><span>Doublons</span></div>
    </div>`;

  const blk=(titre,keys,get,colorOf)=>{
    const rows=keys.map(k=>{
      const set=CARDS.filter(c=>get(c)===k);
      const got=set.filter(c=>has(c.n)).length;
      const p=set.length?got/set.length*100:0;
      const col=colorOf&&colorOf[k] ? colorOf[k] : "var(--rose)";
      return `<div class="row"><span>${k}</span>
        <span class="track" style="color:${col}"><i style="width:${p}%"></i></span>
        <em>${got}/${set.length}</em></div>`;
    }).join("");
    return `<section><h3>${titre}</h3>${rows}</section>`;
  };
  const RCOLOR={"Essence":"#ff4da0","Commune":"#9aa3b2","Peu commune":"#5ac84f","Rare":"#3b82f6",
    "Épique":"#7d4dff","Légendaire":"#e0a82e","Secrète":"#c0304d","Promo":"#22d3ee"};
  const ECOLOR={}; ESSLIST.forEach(k=>{ ECOLOR[k]=`var(${ESS[k]})`; });
  const money = (paid>0||val>0) ? `<section class="money"><h3>Argent</h3>
      <div><b>${eur(paid)}</b><span>investi</span></div>
      <div><b>${eur(val)}</b><span>valeur estimée</span></div>
      <div><b class="${diff>=0?"gain":"loss"}">${diff>=0?"+":""}${eur(diff)}</b><span>écart</span></div>
    </section>` : "";
  document.getElementById("bars").innerHTML=
    blk("Par rareté",RARS,c=>c.rarete,RCOLOR)+
    blk("Par essence",ESSLIST,c=>effEssence(c,store[c.n]),ECOLOR)+
    money;

  const strip=document.getElementById("strip");
  if(strip.childElementCount!==CARDS.length){
    strip.innerHTML=""; CARDS.forEach(c=>{
      const i=document.createElement("i"); i.title=`${c.n} · ${c.nom}`;
      i.onclick=()=>open(c.n); strip.appendChild(i);
    });
  }
  [...strip.children].forEach((i,k)=>{
    const c=CARDS[k];
    i.className = has(c.n)?"has":"";
    i.style.background = has(c.n)?`var(${ESS[effEssence(c,store[c.n])]||"--neutre"})`:"var(--none)";
  });
}

/* ---------- panneau ---------- */
let current=null;
function open(n){
  current=n;
  const c=CARDS.find(x=>x.n===n), s=S(n);
  const p=document.getElementById("panel");
  const ess=effEssence(c,s);
  p.style.setProperty("--essence",`var(${ESS[ess]||"--neutre"})`);
  p.innerHTML=`
    <button class="close" id="pClose" aria-label="Fermer">×</button>
    <div class="pnum">Carte ${c.n} · ${extLogo(c.ext||"Tout Commence Ici","pnum-ext-logo")}</div>
    <h2>${c.nom}</h2>
    <div class="tags">
      <span class="tag">${c.rarete||"Rareté inconnue"}</span>
      <span class="tag">${c.type||"Type inconnu"}</span>
      <select class="tag ess" data-f="essence" aria-label="Essence">
        <option value=""${s.essence?"":" selected"}>${c.essence||"Essence inconnue"}</option>
        ${ESSLIST.map(e=>`<option value="${e}"${s.essence===e?" selected":""}>${e}</option>`).join("")}
      </select>
    </div>
    <div class="preview" id="pPrev"></div>
    <div class="form">
      <section class="pblock">
        <h3>Exemplaires possédés</h3>
        <label>Une version reverse existe&nbsp;?
          <select data-f="rev">
            <option value=""${s.rev===""?" selected":""}>Automatique (${autoRev(c)?"oui":"non"}, selon la rareté)</option>
            <option value="Oui"${s.rev==="Oui"?" selected":""}>Oui, corriger</option>
            <option value="Non"${s.rev==="Non"?" selected":""}>Non, corriger</option>
          </select></label>
        <div id="pOwn"></div>
      </section>
      <section class="pblock">
        <h3>État &amp; valeur</h3>
        <label>État
          <select data-f="etat">${ETATS.map(e=>
            `<option value="${e}"${s.etat===e?" selected":""}>${e||"Non renseigné"}</option>`).join("")}</select></label>
        <div class="pair">
          <label>Prix payé (€)<input type="number" step="0.01" min="0" data-f="prix" value="${s.prix}"></label>
          <label>Valeur estimée (€)<input type="number" step="0.01" min="0" data-f="val" value="${s.val}"></label>
        </div>
      </section>
      <section class="pblock">
        <h3>Provenance &amp; notes</h3>
        <div class="pair">
          <label>Date d'acquisition<input type="date" data-f="date" value="${s.date}"></label>
          <label>Emplacement<input type="text" data-f="loc" value="${s.loc}" placeholder="Classeur 1, page 3"></label>
        </div>
        <label>Notes<textarea data-f="notes" placeholder="Signature, défaut, échange en cours…">${s.notes}</textarea></label>
      </section>
      ${c.note?`<p class="hint">${c.note}</p>`:""}
      ${isExtra(c.n)?`<button class="btn warn" id="delCard">Retirer cette carte de la liste</button>`:""}
    </div>`;
  const prev=document.getElementById("pPrev");
  const img=new Image();
  img.onload=()=>{ prev.innerHTML=""; prev.appendChild(img); };
  img.onerror=()=>{ prev.innerHTML=`<span>Aucune image trouvée pour <b>images-cartes/${c.n}.png</b></span>`; };
  img.src=cardSrc(c.n); img.alt=c.nom;

  function pLine(k,lab){
    const v=+S(n)[k]||0, on=v>0, dup=Math.max(0,v-1);
    if(k==="qr" && !effRev(c,S(n))) return "";
    return `<div class="ownblk">
      <label class="own${on?" on":""}"><input type="checkbox" data-pown="${k}"${on?" checked":""}>
        <span>${lab}</span></label>
      ${on?`<div class="dup"><span>Doublons</span>
        <span class="stepper"><button data-s="${k}" data-d="-1"${dup===0?" disabled":""}>−</button>
        <output>${dup}</output><button data-s="${k}" data-d="1">+</button></span></div>`:""}
    </div>`;
  }
  function paintOwn(){
    const host=document.getElementById("pOwn");
    host.innerHTML=pLine("qn","Normale — celle du classeur")+pLine("qr","Reverse");
    host.querySelectorAll("[data-pown]").forEach(cb=>cb.onchange=ev=>{
      S(n)[cb.dataset.pown]=ev.target.checked?1:0; save(); paintOwn(); render();
    });
    host.querySelectorAll(".stepper button").forEach(b=>b.onclick=()=>{
      S(n)[b.dataset.s]=Math.max(1,(+S(n)[b.dataset.s]||1)+(+b.dataset.d));
      save(); paintOwn(); render();
    });
  }
  paintOwn();

  p.querySelectorAll("[data-f]").forEach(el=>{
    const ev = el.tagName==="SELECT" ? "change" : "input";
    el.addEventListener(ev,()=>{ S(n)[el.dataset.f]=el.value; save(); stats(); });
    el.addEventListener("change",()=>{ paintOwn(); render(); });
  });
  const essSel=p.querySelector('[data-f="essence"]');
  if(essSel) essSel.addEventListener("change",()=>{
    p.style.setProperty("--essence",`var(${ESS[effEssence(c,S(n))]||"--neutre"})`);
  });
  const del=document.getElementById("delCard");
  if(del) del.onclick=()=>{
    if(!confirm("Retirer « "+c.nom+" » de ta liste ? Tes saisies sur cette carte seront perdues.")) return;
    EXTRA=EXTRA.filter(x=>x.n!==c.n); delete store[c.n];
    saveExtra(); save(); close(); render(); toast("Carte retirée");
  };
  document.getElementById("pClose").onclick=close;
  p.classList.add("open"); document.getElementById("scrim").classList.add("open");
}

function nextNum(){
  let max=0;
  CARDS.forEach(c=>{ const m=/^P(\d+)$/.exec(c.n); if(m) max=Math.max(max,+m[1]); });
  return "P"+String(max+1).padStart(2,"0");
}
function openNew(defaultExt, afterSave){
  current=null;
  const ext0 = defaultExt || currentExt || "Tout Commence Ici";
  const p=document.getElementById("panel");
  p.style.setProperty("--essence","var(--neutre)");
  const opts=(arr,sel)=>arr.map(v=>`<option${v===sel?" selected":""}>${v}</option>`).join("");
  p.innerHTML=`
    <button class="close" id="pClose" aria-label="Fermer">×</button>
    <div class="pnum">Nouvelle entrée</div>
    <h2>Ajouter une carte</h2>
    <div class="form">
      <label>Nom de la carte<input id="cNom" type="text" placeholder="Le Marionnettiste - Running the World"></label>
      <div class="pair">
        <label>Numéro<input id="cNum" type="text" value="${nextNum()}"></label>
        <label>Extension<input id="cExt" type="text" value="${ext0}"></label>
      </div>
      <div class="pair">
        <label>Rareté<select id="cRar"><option value="">Non renseignée</option>${opts(RARS,"Promo")}</select></label>
        <label>Type<select id="cTyp"><option value="">Non renseigné</option>${opts(TYPES)}</select></label>
      </div>
      <label>Essence<select id="cEss"><option value="">Non renseignée</option>${opts(ESSLIST)}</select></label>
      <p class="hint">Le visuel apparaîtra si tu déposes une image nommée d'après le numéro, par exemple
        <b>${nextNum()}.png</b>, dans ton dossier images-cartes.</p>
      <div class="pair">
        <button class="btn" id="cSave">Ajouter la carte</button>
        <button class="btn" id="cSaveMore">Ajouter et continuer</button>
      </div>
    </div>`;
  document.getElementById("pClose").onclick=close;
  function doSave(keepOpen){
    const nom=document.getElementById("cNom").value.trim();
    const n=document.getElementById("cNum").value.trim();
    if(!nom){ toast("Il faut au moins un nom"); return; }
    if(!n){ toast("Il faut un numéro"); return; }
    if(CARDS.some(c=>c.n===n)){ toast("Ce numéro est déjà pris"); return; }
    const ext=document.getElementById("cExt").value.trim()||"Tout Commence Ici";
    EXTRA.push({n, nom,
      rarete:document.getElementById("cRar").value,
      type:document.getElementById("cTyp").value,
      essence:document.getElementById("cEss").value,
      ext,
      note:"Carte que tu as ajoutée toi-même"});
    saveExtra();
    toast("« "+nom+" » ajoutée");
    if(keepOpen){ openNew(ext, afterSave); }
    else { close(); if(afterSave) afterSave(); else render(); }
  }
  document.getElementById("cSave").onclick=()=>doSave(false);
  document.getElementById("cSaveMore").onclick=()=>doSave(true);
  p.classList.add("open"); document.getElementById("scrim").classList.add("open");
}
document.getElementById("addCard").onclick=()=>openNew();
function close(){
  current=null;
  document.getElementById("panel").classList.remove("open");
  document.getElementById("scrim").classList.remove("open");
}
document.getElementById("scrim").onclick=close;
document.addEventListener("keydown",e=>{ if(e.key==="Escape") close(); });

/* ---------- import / export ---------- */
function toast(msg){
  const t=document.getElementById("toast");
  t.textContent=msg; t.classList.add("show");
  clearTimeout(t._h); t._h=setTimeout(()=>t.classList.remove("show"),2600);
}
function download(name, text, type){
  const a=document.createElement("a");
  a.href=URL.createObjectURL(new Blob([text],{type}));
  a.download=name; a.click(); URL.revokeObjectURL(a.href);
}
function exportJson(){
  download("ma-collection-m8.json",
    JSON.stringify({v:3,date:new Date().toISOString(),data:store,extra:EXTRA,
      overrides:OVERRIDES,deleted:DELETED,extsExtra:EXT_EXTRA,extMeta:EXT_META},null,1),
    "application/json");
  toast("Fichier exporté");
}
function exportCsv(){
  const head=["N°","Nom","Rareté","Type","Essence","Qté normale","Qté reverse","Reverse existe",
    "État","Prix payé","Valeur estimée","Date","Emplacement","Notes","Statut"];
  const esc=v=>`"${String(v??"").replace(/"/g,'""')}"`;
  const lines=[head.join(";")].concat(CARDS.map(c=>{
    const s=S(c.n);
    return [c.n,c.nom,c.rarete,c.type,effEssence(c,s),+s.qn||0,+s.qr||0,effRev(c,s)?"Oui":"Non",s.etat,s.prix,s.val,
      s.date,s.loc,s.notes,statut(c.n)].map(esc).join(";");
  }));
  download("ma-collection-m8.csv","\ufeff"+lines.join("\r\n"),"text/csv;charset=utf-8");
  toast("CSV exporté");
}
function copyText(text, msg){
  const done=()=>toast(msg);
  const fallback=()=>{
    const ta=document.createElement("textarea");
    ta.value=text; ta.style.position="fixed"; ta.style.opacity="0";
    document.body.appendChild(ta); ta.select();
    try{ document.execCommand("copy"); done(); }catch(e){ toast("Impossible de copier"); }
    ta.remove();
  };
  if(navigator.clipboard && navigator.clipboard.writeText){
    navigator.clipboard.writeText(text).then(done).catch(fallback);
  } else fallback();
}
function copyMissing(){
  const lines=CARDS.filter(c=>statut(c.n)==="manquante").map(c=>`${c.n} · ${c.nom}`);
  if(!lines.length){ toast("Tu as toute la collection !"); return; }
  copyText(lines.join("\n"), `${lines.length} carte(s) manquantes copiées`);
}
function copyDup(){
  const lines=[];
  CARDS.forEach(c=>{
    const s=S(c.n);
    const dupN=Math.max(0,(+s.qn||0)-1), dupR=Math.max(0,(+s.qr||0)-1);
    if(dupN<=0 && dupR<=0) return;
    const parts=[];
    if(dupN>0) parts.push(`${dupN} normale${dupN>1?"s":""}`);
    if(dupR>0) parts.push(`${dupR} reverse${dupR>1?"s":""}`);
    lines.push(`${c.n} · ${c.nom} — ${parts.join(", ")}`);
  });
  if(!lines.length){ toast("Aucun doublon pour l'instant"); return; }
  copyText(lines.join("\n"), `${lines.length} carte(s) en double copiées`);
}
function importPick(){ document.getElementById("file").click(); }
document.getElementById("file").onchange=e=>{
  const f=e.target.files[0]; if(!f) return;
  const r=new FileReader();
  r.onload=()=>{
    try{
      const j=JSON.parse(r.result);
      const d=j.data||j;
      if(typeof d!=="object") throw 0;
      store=d;
      if(Array.isArray(j.extra)){ EXTRA=j.extra; saveExtra(); }
      if(j.overrides && typeof j.overrides==="object"){ OVERRIDES=j.overrides; saveOverrides(); }
      if(Array.isArray(j.deleted)){ DELETED=j.deleted; saveDeleted(); }
      if(Array.isArray(j.extsExtra)){ EXT_EXTRA=j.extsExtra; saveExtExtra(); }
      if(j.extMeta && typeof j.extMeta==="object"){ EXT_META=j.extMeta; saveExtMeta(); }
      save(); render();
      toast("Collection importée");
    }catch(err){ toast("Fichier illisible"); }
    e.target.value="";
  };
  r.readAsText(f);
};
function resetAll(){
  if(!confirm("Effacer toutes tes saisies ? Les cartes du set resteront, mais quantités, prix et notes seront perdus."+(SYNC?" La sauvegarde en ligne sera aussi vidée (l'historique reste sur GitHub).":""))) return;
  store={}; save(); render(); toast("Données effacées");
}

/* ---------- réglages (roue crantée) ---------- */
function openSettings(){
  current=null;
  const p=document.getElementById("panel");
  p.style.setProperty("--essence","var(--neutre)");
  const inApp=!document.getElementById("app").hidden;
  const adm = adminUnlocked ? `
    <section class="pblock"><h3>Administration</h3><div class="setbtns">
      ${inApp?`<button class="btn" id="stAdminCards">Gérer les cartes de cette extension</button>`:""}
      <button class="btn" id="stNewExt">Créer une extension</button>
      <button class="btn" id="stAdminOff">Quitter le mode administration</button>
    </div></section>` : "";
  p.innerHTML=`<button class="close" id="pClose" aria-label="Fermer">×</button>
    <h2>Réglages</h2>
    <div class="form">
      <section class="pblock"><h3>Synchronisation</h3>
        <p class="hint" style="margin:0">${syncText || "Mode local : les saisies restent sur cet appareil"}</p>
        <button class="btn" id="stSync">${SYNC?"Voir la synchronisation":"Synchroniser mes appareils"}</button>
      </section>
      <section class="pblock"><h3>Mes données</h3><div class="setbtns">
        <button class="btn" id="stExJson">Exporter mes données</button>
        <button class="btn" id="stImport">Importer un fichier</button>
        <button class="btn" id="stCsv">Exporter en CSV</button>
      </div>
      <p class="hint" style="margin:0">Sans synchronisation, tes saisies restent dans ce navigateur, sur cet appareil. Exporte-les de temps en temps pour garder une copie.</p>
      </section>
      <section class="pblock"><h3>Partager</h3><div class="setbtns">
        <button class="btn" id="stMiss">Copier les manquantes</button>
        <button class="btn" id="stDup">Copier les doublons</button>
      </div></section>
      ${adm}
      <section class="pblock"><h3>Zone dangereuse</h3>
        <button class="btn warn" id="stReset">Tout effacer</button>
      </section>
    </div>`;
  const on=(id,fn)=>{ const el=document.getElementById(id); if(el) el.onclick=fn; };
  on("pClose",close);
  on("stSync",openSync);
  on("stExJson",exportJson);
  on("stImport",importPick);
  on("stCsv",exportCsv);
  on("stMiss",copyMissing);
  on("stDup",copyDup);
  on("stReset",()=>{ resetAll(); });
  on("stAdminCards",()=>{ close(); setTab("admin"); });
  on("stNewExt",()=>openExtForm(null));
  on("stAdminOff",()=>{ toggleAdmin(); openSettings(); });
  p.classList.add("open"); document.getElementById("scrim").classList.add("open");
}
document.getElementById("gearHome").onclick=openSettings;
document.getElementById("gearApp").onclick=openSettings;

/* ---------- contrôles ---------- */
document.getElementById("q").oninput=e=>{ query=e.target.value; render(); };
document.getElementById("sort").onchange=e=>{ sortBy=e.target.value; render(); };
document.getElementById("vGrid").onclick=()=>{ view="grid";
  document.getElementById("vGrid").setAttribute("aria-pressed","true");
  document.getElementById("vList").setAttribute("aria-pressed","false"); render(); };
document.getElementById("vList").onclick=()=>{ view="list";
  document.getElementById("vList").setAttribute("aria-pressed","true");
  document.getElementById("vGrid").setAttribute("aria-pressed","false"); render(); };

function fActive(){ return fSt.size + fRa.size + fEs.size; }
function updateFCount(){
  const n=fActive(), el=document.getElementById("fCount");
  el.textContent=n; el.classList.toggle("on", n>0);
}
document.getElementById("quickMode").onclick=()=>{
  quick=!quick;
  document.getElementById("quickMode").setAttribute("aria-pressed", quick);
  document.getElementById("quickbar").toggleAttribute("hidden", !quick);
  document.getElementById("out").classList.toggle("quick", quick);
};
document.getElementById("toggleF").onclick=()=>{
  const p=document.getElementById("filters"), open=p.hasAttribute("hidden");
  p.toggleAttribute("hidden", !open);
  document.getElementById("toggleF").setAttribute("aria-expanded", open);
};
document.getElementById("clearF").onclick=()=>{
  fSt.clear(); fRa.clear(); fEs.clear(); buildFilters(); render();
};

/* ---------- administration : cartes et extensions ---------- */
function renderAdmin(){
  const host=document.getElementById("out");
  const list=CARDS.filter(c=>(c.ext||"Tout Commence Ici")===currentExt)
    .sort((a,b)=>a.n.localeCompare(b.n,"fr",{numeric:true}));
  const rows=list.map(c=>`<tr>
    <td class="n">${c.n}</td><td>${c.nom}</td><td>${c.rarete||"—"}</td>
    <td>${c.type||"—"}</td><td>${c.essence||"—"}</td>
    <td><button class="btn" data-edit="${c.n}">Modifier</button>
      <button class="btn warn" data-del="${c.n}">Supprimer</button></td>
  </tr>`).join("");
  host.innerHTML=`<div class="cmpwrap">
    <div class="cmpactions">
      <button class="btn" id="adminAdd">Ajouter une carte à « ${currentExt} »</button>
      <button class="btn" id="adminEditExt">Modifier « ${currentExt} » (info, logo)</button>
      ${currentExt!=="Tout Commence Ici"?`<button class="btn warn" id="adminDeleteExt">Supprimer « ${currentExt} »</button>`:""}
    </div>
    <div class="table-wrap"><table><thead><tr>
      <th>N°</th><th>Nom</th><th>Rareté</th><th>Type</th><th>Essence</th><th>Actions</th>
    </tr></thead><tbody>${rows||`<tr><td colspan="6">Aucune carte dans cette extension pour l'instant.</td></tr>`}</tbody></table></div>
  </div>`;
  document.getElementById("adminAdd").onclick=()=>openNew(currentExt, renderAdmin);
  document.getElementById("adminEditExt").onclick=()=>openExtForm(currentExt);
  const delExtBtn=document.getElementById("adminDeleteExt");
  if(delExtBtn) delExtBtn.onclick=()=>deleteExtension(currentExt);
  host.querySelectorAll("[data-edit]").forEach(b=>b.onclick=()=>openEdit(b.dataset.edit));
  host.querySelectorAll("[data-del]").forEach(b=>b.onclick=()=>adminDelete(b.dataset.del));
}
function adminDelete(n){
  const c=CARDS.find(x=>x.n===n);
  if(!c) return;
  if(!confirm("Supprimer « "+c.nom+" » ("+n+") de la liste ? Tes saisies sur cette carte seront perdues.")) return;
  if(isExtra(n)){ EXTRA=EXTRA.filter(x=>x.n!==n); saveExtra(); }
  else { DELETED.push(n); saveDeleted(); }
  delete store[n]; save();
  renderAdmin(); toast("Carte retirée");
}
function deleteExtension(ext){
  if(ext==="Tout Commence Ici"){ toast("Impossible de supprimer l'extension d'origine"); return; }
  const count=CARDS.filter(c=>(c.ext||"Tout Commence Ici")===ext).length;
  const msg="Supprimer l'extension « "+ext+" »"+(count?" et sa/ses "+count+" carte(s)":"")+" ? Cette action est irréversible.";
  if(!confirm(msg)) return;
  EXTRA.forEach(c=>{ if((c.ext||"Tout Commence Ici")===ext) delete store[c.n]; });
  EXTRA=EXTRA.filter(c=>(c.ext||"Tout Commence Ici")!==ext);
  saveExtra();
  Object.keys(OVERRIDES).forEach(n=>{
    if(OVERRIDES[n].ext===ext){
      delete OVERRIDES[n].ext;
      if(Object.keys(OVERRIDES[n]).length===0) delete OVERRIDES[n];
    }
  });
  saveOverrides();
  save();
  EXT_EXTRA=EXT_EXTRA.filter(e=>e!==ext);
  saveExtExtra();
  delete EXT_META[ext]; saveExtMeta();
  toast("Extension supprimée");
  if(currentExt===ext){
    currentExt=null;
    try{ localStorage.removeItem(EXTKEY); }catch(e){}
    goHome();
  } else {
    renderAdmin();
  }
}
function openEdit(n){
  const c=CARDS.find(x=>x.n===n);
  if(!c) return;
  current=null;
  const p=document.getElementById("panel");
  p.style.setProperty("--essence",`var(${ESS[effEssence(c,store[n])]||"--neutre"})`);
  const opts=(arr,sel)=>arr.map(v=>`<option${v===sel?" selected":""}>${v}</option>`).join("");
  const editableNum=isExtra(n);
  p.innerHTML=`
    <button class="close" id="pClose" aria-label="Fermer">×</button>
    <div class="pnum">Modifier la carte</div>
    <h2>${c.nom}</h2>
    <div class="form">
      <label>Nom de la carte<input id="eNom" type="text" value="${c.nom}"></label>
      <div class="pair">
        <label>Numéro<input id="eNum" type="text" value="${c.n}"${editableNum?"":" disabled"}></label>
        <label>Extension<input id="eExt" type="text" value="${c.ext||"Tout Commence Ici"}"></label>
      </div>
      <div class="pair">
        <label>Rareté<select id="eRar"><option value="">Non renseignée</option>${opts(RARS,c.rarete)}</select></label>
        <label>Type<select id="eTyp"><option value="">Non renseigné</option>${opts(TYPES,c.type)}</select></label>
      </div>
      <label>Essence<select id="eEss"><option value="">Non renseignée</option>${opts(ESSLIST,c.essence)}</select></label>
      ${editableNum?"":`<p class="hint">Cette carte vient du set officiel : son numéro ne peut pas changer, mais le reste oui.</p>`}
      <button class="btn" id="eSave">Enregistrer</button>
    </div>`;
  document.getElementById("pClose").onclick=close;
  document.getElementById("eSave").onclick=()=>{
    const nom=document.getElementById("eNom").value.trim()||c.nom;
    const ext=document.getElementById("eExt").value.trim()||"Tout Commence Ici";
    const patch={nom, ext,
      rarete:document.getElementById("eRar").value,
      type:document.getElementById("eTyp").value,
      essence:document.getElementById("eEss").value};
    if(editableNum){
      const newNum=document.getElementById("eNum").value.trim()||n;
      if(newNum!==n && CARDS.some(x=>x.n===newNum)){ toast("Ce numéro est déjà pris"); return; }
      const idx=EXTRA.findIndex(x=>x.n===n);
      EXTRA[idx]=Object.assign({}, EXTRA[idx], patch, {n:newNum});
      if(newNum!==n && store[n]){ store[newNum]=store[n]; delete store[n]; save(); }
      saveExtra();
    } else {
      OVERRIDES[n]=patch; saveOverrides();
    }
    close(); renderAdmin(); toast("Carte modifiée");
  };
  p.classList.add("open"); document.getElementById("scrim").classList.add("open");
}
function shrinkImage(file, maxDim){
  return new Promise(resolve=>{
    const reader=new FileReader();
    reader.onload=()=>{
      const img=new Image();
      img.onload=()=>{
        let w=img.width, h=img.height;
        if(w>maxDim || h>maxDim){
          const scale=Math.min(maxDim/w, maxDim/h);
          w=Math.round(w*scale); h=Math.round(h*scale);
        }
        const canvas=document.createElement("canvas");
        canvas.width=w; canvas.height=h;
        canvas.getContext("2d").drawImage(img,0,0,w,h);
        resolve(canvas.toDataURL("image/png"));
      };
      img.onerror=()=>resolve(null);
      img.src=reader.result;
    };
    reader.onerror=()=>resolve(null);
    reader.readAsDataURL(file);
  });
}
function openExtForm(existingName){
  const isEdit=!!existingName;
  const meta=isEdit ? Object.assign({info:"",logo:""}, EXT_META[existingName]) : {info:"",logo:""};
  current=null;
  const p=document.getElementById("panel");
  p.style.setProperty("--essence","var(--neutre)");
  p.innerHTML=`
    <button class="close" id="pClose" aria-label="Fermer">×</button>
    <div class="pnum">${isEdit?"Modifier l'extension":"Nouvelle extension"}</div>
    <h2>${isEdit?existingName:"Créer une extension"}</h2>
    <div class="form">
      <label>Nom de l'extension<input id="xName" type="text" value="${isEdit?existingName:""}"
        placeholder="Nom de la prochaine extension" ${isEdit?"disabled":""}></label>
      <label>Info<textarea id="xInfo" placeholder="Une phrase ou deux sur cette extension…">${meta.info}</textarea></label>
      <label>Logo<input id="xLogo" type="file" accept="image/*"></label>
      <div class="logo-preview" id="xPrev">${meta.logo?`<img src="${meta.logo}" alt="">`:`<span>Aucun logo pour l'instant</span>`}</div>
      <button class="btn" id="xSave">${isEdit?"Enregistrer":"Créer l'extension"}</button>
    </div>`;
  document.getElementById("pClose").onclick=close;
  document.getElementById("xLogo").onchange=async e=>{
    const f=e.target.files[0]; if(!f) return;
    const dataUrl=await shrinkImage(f, 700);
    const prev=document.getElementById("xPrev");
    if(dataUrl){ prev.innerHTML=`<img src="${dataUrl}" alt="">`; prev.dataset.pending=dataUrl; }
    else toast("Image illisible");
  };
  document.getElementById("xSave").onclick=()=>{
    const name=isEdit ? existingName : document.getElementById("xName").value.trim();
    if(!name){ toast("Il faut un nom"); return; }
    if(!isEdit && extsList().includes(name)){ toast("Cette extension existe déjà"); return; }
    const pending=document.getElementById("xPrev").dataset.pending;
    EXT_META[name]={ info:document.getElementById("xInfo").value.trim(), logo: pending || meta.logo || "" };
    saveExtMeta();
    close();
    if(isEdit){ renderAdmin(); toast("Extension mise à jour"); }
    else{
      EXT_EXTRA.push(name); saveExtExtra();
      enterExt(name); setTab("admin"); toast("« "+name+" » créée");
    }
  };
  p.classList.add("open"); document.getElementById("scrim").classList.add("open");
}

/* ---------- comparaison avec des amis ---------- */
let cmpFriend=0, cmpFilter="all", cmpMore={give:false,ask:false};
const cmpEsc=s=>String(s==null?"":s).replace(/[&<>"]/g,ch=>({"&":"&amp;","<":"&lt;",">":"&gt;",'"':"&quot;"}[ch]));
function qOf(data,n){ const s=data[n]||{}; return [+s.qn||0, +s.qr||0]; }
/* cartes que `from` peut donner à `to` : doublons que l'autre n'a pas du tout dans cette version */
function tradeList(from,to){
  const out=[];
  CARDS.forEach(c=>{
    const [fn,fr]=qOf(from,c.n), [tn,tr]=qOf(to,c.n);
    const parts=[];
    if(fn>1 && tn===0) parts.push(`en double ×${fn-1}`);
    if(fr>1 && tr===0 && effRev(c,from[c.n])) parts.push(`reverse en double ×${fr-1}`);
    if(parts.length) out.push({c, txt:parts.join(" · ")});
  });
  return out;
}
function cellHTML(data,c){
  const [n,r]=qOf(data,c.n);
  const dup=Math.max(0,n-1)+Math.max(0,r-1);
  const {cls,txt}=versionState(data,c);
  return `<span class="st ${cls}">${txt}</span>${dup>0?`<span class="stdup">${dup} en double</span>`:""}`;
}
function tradeText(items){ return items.map(it=>`${it.c.n} · ${it.c.nom} — ${it.txt}`).join("\n"); }
function renderCompare(){
  const host=document.getElementById("out");
  if(!FRIENDS.length){
    host.innerHTML=`<div class="empty"><b>Aucun ami importé</b>
      Demande à tes amis d'exporter leur collection (bouton « Exporter mes données » dans la roue crantée),
      puis importe leur(s) fichier(s) ici pour comparer vos collections.<br><br>
      <button class="btn primary" id="addFriends">Importer des amis</button></div>`;
    document.getElementById("addFriends").onclick=()=>document.getElementById("friendFile").click();
    return;
  }
  if(cmpFriend>=FRIENDS.length || cmpFriend<0) cmpFriend=0;
  const f=FRIENDS[cmpFriend], nm=cmpEsc(f.name);
  const give=tradeList(store,f.data), ask=tradeList(f.data,store);
  const theirN=ownedCount(f.data);
  const common=CARDS.filter(c=>{ const [a,b]=qOf(f.data,c.n); return has(c.n) && a+b>0; }).length;

  const listHTML=(key,title,items)=>{
    const limit=8, all=cmpMore[key], shown=all?items:items.slice(0,limit);
    return `<section class="xbox"><div class="xhead"><b>${title} · ${items.length}</b>
        ${items.length?`<button class="btn small" data-copy="${key}">Copier</button>`:""}</div>
      ${items.length
        ? shown.map(it=>`<div class="xrow"><span><span class="n">${it.c.n}</span>${cmpEsc(it.c.nom)}</span><em>${it.txt}</em></div>`).join("")
        : `<p class="xempty">Rien à échanger dans ce sens pour l'instant.</p>`}
      ${items.length>limit?`<button class="xmore" data-more="${key}">${all?"Réduire":`Voir ${items.length-limit} autre${items.length-limit>1?"s":""}`}</button>`:""}
    </section>`;
  };

  const rows=CARDS.filter(c=>{
    const [mn,mr]=qOf(store,c.n), [tn,tr]=qOf(f.data,c.n);
    const iHave=mn+mr>0, heHas=tn+tr>0;
    if(cmpFilter==="they") return heHas && !iHave;
    if(cmpFilter==="me") return iHave && !heHas;
    return true;
  }).map(c=>`<tr><td class="n">${c.n}</td><td>${cmpEsc(c.nom)}</td>
      <td>${cellHTML(store,c)}</td><td>${cellHTML(f.data,c)}</td></tr>`).join("");

  const flt=[["all","Toutes les cartes"],["they",`Chez ${nm}, pas chez toi`],["me",`Chez toi, pas chez ${nm}`]];
  host.innerHTML=`<div class="cmpwrap">
    <div class="xchips">
      ${FRIENDS.map((x,i)=>`<button class="chip" data-f="${i}" aria-pressed="${i===cmpFriend}">${cmpEsc(x.name)}</button>`).join("")}
      <button class="chip" id="addFriends">+ Importer un ami</button>
    </div>
    <div class="minis">
      <div class="stat"><b>${theirN}</b><span>cartes chez ${nm}</span></div>
      <div class="stat"><b>${common}</b><span>cartes en commun</span></div>
      <div class="stat"><b>${give.length} ⇄ ${ask.length}</b><span>échanges possibles</span></div>
    </div>
    <div class="xlists">
      ${listHTML("give",`Tu peux donner à ${nm}`,give)}
      ${listHTML("ask",`${nm} peut te donner`,ask)}
    </div>
    <section class="xbox">
      <div class="xhead"><b>Toutes les cartes</b></div>
      <div class="xfilters">${flt.map(([k,l])=>`<button class="chip" data-flt="${k}" aria-pressed="${cmpFilter===k}">${l}</button>`).join("")}</div>
      <div class="table-wrap" style="margin-top:10px"><table class="xtable">
        <thead><tr><th>N°</th><th>Nom</th><th>Toi</th><th>${nm}</th></tr></thead>
        <tbody>${rows||`<tr><td colspan="4">Aucune carte dans cette liste.</td></tr>`}</tbody>
      </table></div>
    </section>
    <div class="cmpactions">
      <button class="btn" id="copyTrade">Copier le résumé d'échange</button>
      <button class="btn warn" id="removeFriend">Retirer ${nm}</button>
    </div>
  </div>`;

  host.querySelectorAll("[data-f]").forEach(b=>b.onclick=()=>{ cmpFriend=+b.dataset.f; cmpMore={give:false,ask:false}; renderCompare(); });
  host.querySelectorAll("[data-flt]").forEach(b=>b.onclick=()=>{ cmpFilter=b.dataset.flt; renderCompare(); });
  host.querySelectorAll("[data-more]").forEach(b=>b.onclick=()=>{ cmpMore[b.dataset.more]=!cmpMore[b.dataset.more]; renderCompare(); });
  host.querySelectorAll("[data-copy]").forEach(b=>b.onclick=()=>{
    const k=b.dataset.copy, items=k==="give"?give:ask;
    copyText(tradeText(items), `${items.length} carte(s) copiée(s)`);
  });
  document.getElementById("addFriends").onclick=()=>document.getElementById("friendFile").click();
  document.getElementById("copyTrade").onclick=()=>{
    const lines=[`Échange avec ${f.name} :`,"",
      `Je peux te donner (${give.length}) :`, ...(give.length?[tradeText(give)]:["—"]), "",
      `Tu peux me donner (${ask.length}) :`, ...(ask.length?[tradeText(ask)]:["—"])];
    copyText(lines.join("\n"), `Résumé avec ${f.name} copié`);
  };
  document.getElementById("removeFriend").onclick=()=>{
    FRIENDS.splice(cmpFriend,1); cmpFriend=0; cmpMore={give:false,ask:false}; renderCompare();
  };
}
document.getElementById("friendFile").onchange=e=>{
  const files=[...e.target.files]; if(!files.length) return;
  let done=0; const firstNew=FRIENDS.length;
  files.forEach(f=>{
    const r=new FileReader();
    r.onload=()=>{
      try{
        const j=JSON.parse(r.result);
        const d=j.data||j;
        if(typeof d!=="object" || Array.isArray(d)) throw 0;
        FRIENDS.push({name:f.name.replace(/\.json$/i,""), data:d});
      }catch(err){ toast("Fichier illisible : "+f.name); }
      done++;
      if(done===files.length){
        cmpFriend=Math.max(0,Math.min(firstNew,FRIENDS.length-1)); cmpMore={give:false,ask:false};
        setTab("compare"); toast(FRIENDS.length+" ami(s) chargé(s)");
      }
    };
    r.readAsText(f);
  });
  e.target.value="";
};

/* ---------- onglets ---------- */
function setTab(t){
  const stats = t==="stats", compare = t==="compare", admin = t==="admin";
  const away = stats||compare||admin;
  document.getElementById("board").style.display = stats ? "grid" : "none";
  document.getElementById("bars").style.display  = stats ? "grid" : "none";
  document.querySelector(".tools").style.display = away ? "none" : "grid";
  document.getElementById("filters").style.display = away ? "none" : "";
  document.getElementById("quickbar").toggleAttribute("hidden", !(quick && !away));
  document.getElementById("out").style.display  = stats ? "none" : "block";
  document.getElementById("tColl").setAttribute("aria-pressed", !away);
  document.getElementById("tStats").setAttribute("aria-pressed", stats);
  document.getElementById("tCompare").setAttribute("aria-pressed", compare);
  document.getElementById("tAdmin").setAttribute("aria-pressed", admin);
  document.getElementById("tAdmin").hidden = !admin;
  if(compare) renderCompare(); else if(admin) renderAdmin(); else if(!stats) render();
}
document.getElementById("tColl").onclick=()=>setTab("coll");
document.getElementById("tStats").onclick=()=>setTab("stats");
document.getElementById("tCompare").onclick=()=>setTab("compare");
document.getElementById("tAdmin").onclick=()=>setTab("admin");
setTab("coll");

const UIKEY="m8tcg.ui";
function setSize(w, persist){
  document.documentElement.style.setProperty("--cardmin", w+"px");
  document.getElementById("sizeRange").value=w;
  if(persist){ try{ localStorage.setItem(UIKEY, w); }catch(e){} }
}
document.getElementById("sizeRange").oninput=e=>setSize(+e.target.value, true);
let savedW = 230;
try{ savedW = +localStorage.getItem(UIKEY) || 230; }catch(e){}
setSize(savedW, false);

document.getElementById("backHome").onclick=goHome;
buildFilters(); loadBrand();
if(SYNC) syncNow(); else setSync("Mode local : les saisies restent sur cet appareil");
if(memoryOnly) toast("Ce navigateur bloque la sauvegarde — exporte tes données avant de fermer");

let savedExt=null;
try{ savedExt=localStorage.getItem(EXTKEY); }catch(e){}
applyAdminVisibility();
armSecretUnlock(document.getElementById("brand"));
armSecretUnlock(document.getElementById("homeBrand"));
armSecretUnlock(document.getElementById("homeTitle"));
armSecretUnlock(document.getElementById("title"));
if(savedExt && extsList().includes(savedExt)) enterExt(savedExt); else goHome();
</script>
</body>
</html>
"""

HTML = HTML.replace("__DATA__", DATA).replace("__TOTAL__", str(len(cards)))
TEMPLATE = HTML                       # gabarit avec __IMAGES__ encore en place
APP = HTML.replace("__IMAGES__", "{}")

with open(os.path.join(BASE_DIR, "ma-collection.html"), "w", encoding="utf-8") as f:
    f.write(APP)
# GitHub Pages sert index.html à la racine du dépôt : même contenu, mêmes chemins relatifs
with open(os.path.join(BASE_DIR, "index.html"), "w", encoding="utf-8") as f:
    f.write(APP)
print("html ok")
